import json
import tempfile
import threading
import unittest
from pathlib import Path

import fitz
from docx import Document
from pydantic import ValidationError

from app.checks.evidence_tools import execute_tool
from app.contracts.agent import AgentOptions, Candidate, strict_schema
from app.documents.extraction import extract_document
from app.documents.extraction.common import DocumentReadCanceled
from app.documents.index import DocumentIndex, text_blocks


class DocumentIndexTest(unittest.TestCase):
    def test_long_paragraph_preserves_sentences_and_exact_original_text(self):
        sentences = [
            f"条件{i}：" + "甲" * 480 + "；禁止带电操作。\n" for i in range(20)
        ]
        original = "".join(sentences).strip()
        index = DocumentIndex([{"filename": "a.txt", "blocks": [{"text": original}]}])
        self.assertEqual("".join(b.text for b in index.blocks), original)
        for sentence in sentences:
            self.assertTrue(any(sentence.strip() in b.text for b in index.blocks))
        self.assertTrue(all(len(b.text) <= 4000 for b in index.blocks))

    def test_long_pdf_table_preserves_rows_and_provides_original_header_context(self):
        rows = ["<tr><td>型号</td><td>额定电压 / V</td></tr>\n"] + [
            f'<tr><td data-cell="A{i}">设备{i}</td><td>{"甲" * 180}</td></tr>\n'
            for i in range(70)
        ]
        text = "<table>\n" + "".join(rows) + "</table>"
        index = DocumentIndex(
            [
                {
                    "filename": "a.pdf",
                    "blocks": [
                        {"text": "仅适用于室内设备。", "section": "设备A"},
                        {
                            "kind": "table",
                            "text": text,
                            "section": "设备A",
                            "location": "第1页 / 表格 T1",
                        },
                        {"text": "注：低温时参数应降额。", "section": "设备A"},
                    ],
                }
            ]
        )
        fragments = [b for b in index.blocks if b.kind == "table"]
        self.assertGreater(len(fragments), 2)
        self.assertEqual("".join(b.text for b in fragments), text)
        for row in rows:
            self.assertTrue(any(row in b.text for b in fragments))
        target = fragments[-1]
        context = index.supporting_sources([target.id])
        self.assertNotIn(target.id, {b["id"] for b in context})
        self.assertTrue(any("额定电压 / V" in b["text"] for b in context))
        self.assertTrue(any("室内设备" in b["text"] for b in context))
        self.assertTrue(any("低温" in b["text"] for b in context))
        self.assertTrue(
            all(
                index.verify({"source_id": b["id"], "quote": b["text"]})
                for b in context
            )
        )

    def test_batches_keep_fitting_sections_and_tables_together(self):
        index = DocumentIndex(
            [
                {
                    "filename": "a.txt",
                    "blocks": [
                        {"text": "甲" * 450, "section": "A"},
                        {"text": "乙" * 300, "section": "B"},
                        {"text": "丙" * 300, "section": "B"},
                        {"text": "丁" * 450, "section": "C"},
                        *[
                            {
                                "text": f"型号{i} | 电压",
                                "kind": "table",
                                "location": f"表格1 > 第{i}行",
                                "section": "C",
                            }
                            for i in range(3)
                        ],
                        {"text": "戊" * 900, "section": "C"},
                    ],
                }
            ]
        )
        chunks = index.chunks(1300)
        positions = {b.id: i for i, chunk in enumerate(chunks) for b in chunk}
        self.assertEqual(positions[index.blocks[1].id], positions[index.blocks[2].id])
        self.assertEqual(
            len({positions[b.id] for b in index.blocks if b.kind == "table"}), 1
        )
        self.assertEqual(
            [b.id for chunk in chunks for b in chunk], [b.id for b in index.blocks]
        )
        self.assertTrue(
            all(
                sum(len(json.dumps(b.as_dict(), ensure_ascii=False)) + 2 for b in chunk)
                <= 1300
                for chunk in chunks
            )
        )

    def test_supporting_context_keeps_table_and_document_identity(self):
        documents = [
            {
                "filename": "same.docx",
                "blocks": [
                    {"text": f"设备{number}适用范围", "section": "参数"},
                    *[
                        {
                            "kind": "table",
                            "text": f"设备{number}参数{i}",
                            "section": "参数",
                            "location": f"表格1 > 第{i + 1}行",
                        }
                        for i in range(12)
                    ],
                    {"text": f"设备{number}注释", "section": "参数"},
                    {"text": "另一章节", "section": "安装"},
                ],
            }
            for number in range(2)
        ]
        index = DocumentIndex(documents)
        target = index.blocks[10]
        context = index.supporting_sources([target.id])
        self.assertTrue(any("设备0参数0" in b["text"] for b in context))
        self.assertTrue(any("设备0注释" in b["text"] for b in context))
        self.assertFalse(
            any("设备1" in b["text"] or b["section"] == "安装" for b in context)
        )

    def test_docx_preserves_paragraph_table_context_and_empty_columns(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "guide.docx"
            doc = Document()
            doc.add_heading("A 型号", 1)
            doc.add_paragraph("下表只适用于 A 型号。")
            table = doc.add_table(rows=1, cols=3)
            table.cell(0, 0).text = "电压"
            table.cell(0, 2).text = "220 V"
            doc.add_heading("B 型号", 1)
            doc.add_paragraph("110 V")
            doc.save(path)
            blocks = []
            text, _ = extract_document(path, "docx", blocks=blocks)
            self.assertLess(text.index("电压"), text.index("B 型号"))
            self.assertIn("电压 |  | 220 V", text)
            source = next(b for b in blocks if b["kind"] == "table")
            self.assertEqual(source["section"], "A 型号")
            self.assertIn("表格1", source["location"])

    def test_docx_cancellation_checks_content_boundaries(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "guide.docx"
            Document().save(path)
            canceled = threading.Event()
            canceled.set()
            with self.assertRaises(DocumentReadCanceled):
                extract_document(path, "docx", cancel_event=canceled)

    def test_docx_headers_footers_and_their_tables_keep_source_locations(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "guide.docx"
            doc = Document()
            doc.add_paragraph("安装指南正文")
            section = doc.sections[0]
            section.header.paragraphs[0].text = "A型号 / 版本2"
            table = section.footer.add_table(rows=1, cols=2, width=section.page_width)
            table.cell(0, 0).text = "发布日期"
            table.cell(0, 1).text = "2026-01-01"
            section.different_first_page_header_footer = True
            section.first_page_header.paragraphs[0].text = "首页说明"
            doc.save(path)
            blocks = []
            text, _ = extract_document(path, "docx", blocks=blocks)
            self.assertIn("首页说明", text)
            self.assertIn("发布日期 | 2026-01-01", text)
            footer = next(b for b in blocks if b["kind"] == "table")
            self.assertIn("页脚 / 表格", footer["location"])
            self.assertEqual(footer["section"], "第1节 / 页脚")

    def test_pdf_records_unread_pages(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "guide.pdf"
            doc = fitz.open()
            doc.new_page().insert_text((72, 72), "Voltage 220 V")
            doc.new_page()
            doc.save(path)
            doc.close()
            blocks = []
            extract_document(path, "pdf", blocks=blocks)
            self.assertEqual(len(blocks), 2)
            self.assertEqual(blocks[1]["kind"], "unresolved")

    def test_reference_requires_exact_quote_and_allowed_source(self):
        index = DocumentIndex(
            [{"filename": "a.txt", "text": "A: 220 V\nB: 110 V", "text_complete": True}]
        )
        first, second = index.blocks
        self.assertTrue(index.verify({"source_id": first.id, "quote": "220 V"}))
        self.assertFalse(index.verify({"source_id": first.id, "quote": "110 V"}))
        self.assertFalse(
            index.verify({"source_id": first.id, "quote": "220 V"}, allowed={second.id})
        )
        self.assertFalse(index.verify({"source_id": first.id, "quote": ""}))

    def test_source_ids_stable_and_documents_separated(self):
        documents = [
            {"filename": "a", "text": "same"},
            {"filename": "b", "text": "same"},
        ]
        index = DocumentIndex(documents)
        self.assertEqual(index.signature, DocumentIndex(documents).signature)
        self.assertNotEqual(index.blocks[0].id, index.blocks[1].id)
        self.assertEqual(len(index.context(index.blocks[0].id, 3)), 1)

    def test_duplicate_filenames_do_not_mix_context_or_table_rows(self):
        index = DocumentIndex(
            [
                {
                    "filename": "guide.docx",
                    "role": "master",
                    "blocks": [
                        {"kind": "table", "location": "表格1 > 第1行", "text": text}
                    ],
                    "text_complete": True,
                }
                for text in ("A型号 | 220 V", "B型号 | 110 V")
            ]
        )
        source = index.blocks[0]
        self.assertEqual(
            [b["text"] for b in index.context(source.id, 3)], [source.text]
        )
        result = execute_tool("read_table", {"source_id": source.id}, index, {})
        self.assertEqual([b["text"] for b in result["sources"]], [source.text])

    def test_table_read_keeps_requested_distant_row_and_headers(self):
        index = DocumentIndex(
            [
                {
                    "filename": "guide.docx",
                    "text_complete": True,
                    "blocks": [
                        {
                            "kind": "table",
                            "location": f"表格1 > 第{i + 1}行",
                            "text": f"型号{i} | 参数{i}",
                        }
                        for i in range(80)
                    ],
                }
            ]
        )
        target = index.blocks[65]
        result = execute_tool("read_table", {"source_id": target.id}, index, {})
        self.assertEqual(result["sources"][0]["id"], target.id)
        self.assertIn(index.blocks[0].id, {b["id"] for b in result["sources"]})
        self.assertEqual(len(result["sources"]), 20)
        self.assertTrue(result["truncated"])

    def test_missing_hits_do_not_establish_complete_coverage(self):
        index = DocumentIndex(
            [{"filename": "a", "text": "first\nsecond", "text_complete": False}]
        )
        self.assertEqual(index.search("missing"), [])
        coverage = index.coverage([b.id for b in index.blocks])
        self.assertTrue(coverage["all_checked"])
        self.assertFalse(coverage["text_complete"])

    def test_text_tables_stay_with_page_context(self):
        blocks = text_blocks(
            "[第2页]\n# Scope\n[PDF结构化表格 T1 开始；x]\n<table>220</table>\n[PDF结构化表格 T1 结束]\nEnd"
        )
        table = next(b for b in blocks if b["kind"] == "table")
        self.assertIn("第2页", table["location"])
        self.assertEqual(table["section"], "Scope")
        self.assertEqual(blocks[-1]["text"], "End")

    def test_candidate_schema_rejects_unrecognized_fields_and_types(self):
        with self.assertRaises(ValidationError):
            Candidate.model_validate(
                {"description": "无证据", "severity": "fatal", "evidence": "none"}
            )
        schema = strict_schema(Candidate)
        self.assertFalse(schema["additionalProperties"])
        self.assertEqual(set(schema["required"]), set(schema["properties"]))
        for invalid_price in (float("nan"), float("inf"), -1.0):
            with self.assertRaises(ValidationError):
                AgentOptions(input_price=invalid_price)


if __name__ == "__main__":
    unittest.main()
