import json
import subprocess
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import fitz
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

from app.documents.extraction import (
    allowed_file,
    extract_document,
    extract_text,
    format_document_text,
)
from app.documents.extraction.pdf import (
    _extract_pymupdf_page_with_tables,
    _pdf_should_try_pypdf,
    _select_pdf_page_text,
)
from app.documents.images import (
    candidate_pdf_pages_for_image_check,
    extract_images,
    format_image_document_text,
    image_items_from_meta,
    image_path_from_item,
    select_pdf_page_numbers,
)
from app.documents.videos import (
    VideoFrameExtractionError,
    _decode_process_output,
    _extract_frame,
    _probe_video_duration,
    _sample_video_timestamps,
    _VideoFrameCommandError,
    allowed_video_file,
    extract_video_frames,
    format_video_document_text,
    video_extension_of,
)

_TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


class DocumentFormattingTest(unittest.TestCase):
    def test_formats_document_text_with_filename(self):
        self.assertEqual(
            format_document_text("../报告.txt", " 内容 "),
            "file: 报告.txt\n\n内容",
        )

    def test_empty_text_stays_empty(self):
        self.assertEqual(format_document_text("报告.txt", "  "), "")

    def test_allows_excel_files(self):
        self.assertTrue(allowed_file("素材.xlsx"))
        self.assertTrue(allowed_file("素材.xlsm"))
        self.assertTrue(allowed_file("素材.xls"))

    def test_extracts_docx_hyperlink_targets_from_paragraph_and_table(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "指南.docx"
            document = Document()
            paragraph = document.add_paragraph("操作步骤请参见")
            _add_docx_hyperlink(
                paragraph,
                "安装指南",
                "https://docs.example.com/install",
            )
            table = document.add_table(rows=1, cols=1)
            cell_paragraph = table.cell(0, 0).paragraphs[0]
            _add_docx_hyperlink(
                cell_paragraph,
                "维护指南",
                "https://docs.example.com/maintenance",
            )
            document.save(path)

            text, hyperlinks = extract_document(path, "docx")

        self.assertIn(
            "操作步骤请参见安装指南（超链接：https://docs.example.com/install）",
            text,
        )
        self.assertIn(
            "维护指南（超链接：https://docs.example.com/maintenance）",
            text,
        )
        self.assertEqual(len(hyperlinks), 2)
        self.assertEqual(hyperlinks[0]["display_text"], "安装指南")
        self.assertEqual(hyperlinks[0]["target"], "https://docs.example.com/install")
        self.assertEqual(hyperlinks[0]["location"], "第1段")

    def test_extracts_xlsx_workbook_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "素材.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.title = "参数表"
            sheet.append(["项目", "参数", "单位"])
            sheet.append(["额定电流", 10, "A"])
            sheet.append(["空列示例", None, "保留位置"])
            sheet["D5"] = "=SUM(1,2)"
            sheet["A6"] = "安装指南"
            sheet["A6"].hyperlink = "https://docs.example.com/install"
            workbook.save(path)
            workbook.close()

            text, hyperlinks = extract_document(path, "xlsx")

        self.assertIn("# 工作表：参数表", text)
        self.assertIn("项目 | 参数 | 单位", text)
        self.assertIn("额定电流 | 10 | A", text)
        self.assertIn("空列示例 |  | 保留位置", text)
        self.assertIn(" |  |  | =SUM(1,2)", text)
        self.assertIn(
            "安装指南（超链接：https://docs.example.com/install）",
            text,
        )
        self.assertEqual(hyperlinks[0]["location"], "工作表“参数表”!A6")
        self.assertEqual(hyperlinks[0]["target"], "https://docs.example.com/install")

    def test_extracts_html_hyperlink_target(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "指南.html"
            path.write_text(
                '<p>具体步骤参见<a href="https://docs.example.com/guide">操作指南</a>。</p>',
                encoding="utf-8",
            )

            text, hyperlinks = extract_document(path, "html")

        self.assertIn(
            "操作指南（超链接：https://docs.example.com/guide）",
            text,
        )
        self.assertEqual(hyperlinks[0]["target"], "https://docs.example.com/guide")

    def test_extracts_html_internal_link_state_and_unsafe_scheme(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "links.html"
            path.write_text(
                """<a href="#existing">已有位置</a>
<a href="#missing">缺失位置</a>
<a href="javascript:alert(1)">不安全链接</a>
<h2 id="existing">目标章节</h2>""",
                encoding="utf-8",
            )

            _text, hyperlinks = extract_document(path, "html")

        self.assertTrue(hyperlinks[0]["internal_target_exists"])
        self.assertFalse(hyperlinks[1]["internal_target_exists"])
        self.assertEqual(hyperlinks[2]["target"], "javascript:alert(1)")

    def test_extracts_markdown_hyperlinks_without_changing_source_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "guide.md"
            source = "请参见[安装指南](https://docs.example.com/install)。"
            path.write_text(source, encoding="utf-8")

            text, hyperlinks = extract_document(path, "md")

        self.assertEqual(text, source)
        self.assertEqual(hyperlinks[0]["display_text"], "安装指南")
        self.assertEqual(hyperlinks[0]["target"], "https://docs.example.com/install")

    def test_pdf_extraction_removes_overlapping_space_but_keeps_visible_space(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "warning.pdf"
            document = fitz.open()
            page = document.new_page()
            font_size = 12
            start_x = 72
            page.insert_text((start_x, 72), "D ", fontsize=font_size)
            page.insert_text(
                (
                    start_x
                    + fitz.get_text_length("D", fontname="helv", fontsize=font_size),
                    72,
                ),
                "ANGER",
                fontsize=font_size,
            )
            page.insert_text((start_x, 100), "NORMAL SPACE", fontsize=font_size)
            document.save(path)
            document.close()

            text = extract_text(path, "pdf")

        self.assertIn("DANGER", text)
        self.assertNotIn("D ANGER", text)
        self.assertIn("NORMAL SPACE", text)

    def test_extracts_pdf_link_annotation_target(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "guide.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Installation Guide", fontsize=12)
            page.insert_link(
                {
                    "kind": fitz.LINK_URI,
                    "from": fitz.Rect(72, 58, 170, 76),
                    "uri": "https://docs.example.com/install",
                }
            )
            document.save(path)
            document.close()

            text, hyperlinks = extract_document(path, "pdf")

        self.assertIn("[超链接]", text)
        self.assertIn("https://docs.example.com/install", text)
        self.assertEqual(hyperlinks[0]["location"], "第1页")
        self.assertEqual(hyperlinks[0]["target"], "https://docs.example.com/install")

    def test_pdf_extraction_skips_pypdf_for_clean_pymupdf_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "clean.pdf"
            document = fitz.open()
            page = document.new_page()
            page.insert_text((72, 72), "Clean PDF text")
            document.save(path)
            document.close()

            with patch(
                "app.documents.extraction.pdf.PdfReader",
                side_effect=AssertionError("不应加载 pypdf"),
            ):
                text = extract_text(path, "pdf")

        self.assertIn("Clean PDF text", text)

    def test_pdf_text_extraction_can_skip_table_structure(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table.pdf"
            document = fitz.open()
            page = document.new_page()
            for x in (50, 150, 250):
                page.draw_line((x, 50), (x, 150))
            for y in (50, 100, 150):
                page.draw_line((50, y), (250, y))
            page.insert_text((65, 80), "A")
            document.save(path)
            document.close()

            with patch(
                "app.documents.extraction.pdf._extract_pymupdf_page_with_tables"
            ) as table_extractor:
                text = extract_text(path, "pdf", include_tables=False)

        self.assertIn("A", text)
        table_extractor.assert_not_called()

    def test_pdf_fallback_predicate_only_accepts_empty_or_corrupted_text(self):
        self.assertTrue(_pdf_should_try_pypdf(""))
        self.assertTrue(_pdf_should_try_pypdf("标题 \ufffd 正文"))
        self.assertFalse(_pdf_should_try_pypdf("标题 完整正文"))

    def test_extracts_pdf_table_in_reading_order_with_merged_and_empty_cells(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "merged-table.pdf"
            document = fitz.open()
            page = document.new_page(width=420, height=300)
            page.insert_text((50, 60), "Before table")
            lines = (
                (50, 100, 350, 100),
                (50, 140, 350, 140),
                (150, 180, 350, 180),
                (50, 220, 350, 220),
                (50, 100, 50, 220),
                (350, 100, 350, 220),
                (150, 140, 150, 220),
                (250, 100, 250, 220),
            )
            for x0, y0, x1, y1 in lines:
                page.draw_line(fitz.Point(x0, y0), fitz.Point(x1, y1))
            for x, y, value in (
                (100, 125, "Merged header"),
                (285, 125, "Column C"),
                (80, 165, "Merged group"),
                (185, 165, "Value 1"),
                (285, 165, "Value 2"),
                (185, 205, "Value 3"),
            ):
                page.insert_text((x, y), value)
            page.insert_text((50, 260), "After table")
            document.save(path)
            document.close()

            text = extract_text(path, "pdf")

        self.assertLess(text.index("Before table"), text.index("[PDF结构化表格"))
        self.assertLess(text.index("[PDF结构化表格"), text.index("After table"))
        self.assertIn(
            '<table id="page001-table001" data-confidence="medium" data-view="normalized">',
            text,
        )
        self.assertIn(
            '<td data-cell="A1" data-original-range="A1:B1" '
            'data-original-colspan="2">Merged header</td>',
            text,
        )
        self.assertIn(
            '<td data-cell="B1" data-original-range="A1:B1" '
            'data-inherited-from="A1">Merged header</td>',
            text,
        )
        self.assertIn(
            '<td data-cell="A2" data-original-range="A2:A3" '
            'data-original-rowspan="2">Merged group</td>',
            text,
        )
        self.assertIn(
            '<td data-cell="A3" data-original-range="A2:A3" '
            'data-inherited-from="A2">Merged group</td>',
            text,
        )
        self.assertIn(
            '<td data-cell="C3" data-empty="true">[空单元格]</td>',
            text,
        )
        self.assertEqual(text.count("Merged header"), 2)
        self.assertEqual(text.count("Merged group"), 2)

    def test_distinguishes_nontext_pdf_table_cell_from_empty_cell(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "table-with-icon.pdf"
            document = fitz.open()
            page = document.new_page(width=300, height=200)
            for x0, y0, x1, y1 in (
                (50, 50, 250, 50),
                (50, 100, 250, 100),
                (50, 150, 250, 150),
                (50, 50, 50, 150),
                (150, 50, 150, 150),
                (250, 50, 250, 150),
            ):
                page.draw_line(fitz.Point(x0, y0), fitz.Point(x1, y1))
            page.insert_text((80, 80), "Icon")
            page.insert_text((180, 80), "Description")
            page.insert_image(fitz.Rect(80, 115, 100, 135), stream=_TINY_PNG)
            document.save(path)
            document.close()

            text = extract_text(path, "pdf")

        self.assertIn(
            '<td data-cell="A2" data-non-text="true">[非文本图形或图标]</td>',
            text,
        )
        self.assertIn(
            '<td data-cell="B2" data-empty="true">[空单元格]</td>',
            text,
        )

    def test_pdf_table_extraction_skips_one_malformed_candidate(self):
        page = MagicMock()
        page.find_tables.return_value.tables = [object(), object()]
        page.get_drawings.return_value = [
            {"items": [("l",), ("l",), ("l",), ("l",)], "rect": (0, 0, 1, 1)}
        ]
        page.get_image_info.return_value = []
        valid_model = {
            "id": "page001-table002",
            "page_number": 1,
            "table_index": 2,
            "bbox": (50.0, 50.0, 250.0, 150.0),
            "row_count": 2,
            "column_count": 2,
            "cells": [
                {
                    "row": 0,
                    "column": 0,
                    "rowspan": 1,
                    "colspan": 1,
                    "bbox": (50.0, 50.0, 150.0, 100.0),
                    "text": "项目",
                    "has_nontext_content": False,
                }
            ],
            "confidence": "high",
            "unresolved_positions": [],
            "vector_edge_count": 6,
        }

        with patch(
            "app.documents.extraction.pdf._pdf_table_model",
            side_effect=(RuntimeError("malformed table"), valid_model),
        ):
            text = _extract_pymupdf_page_with_tables(page, 1, {}, "")

        self.assertIn("page001-table002", text)
        self.assertIn("项目", text)

    def test_pdf_table_extraction_skips_pages_without_vector_edges(self):
        page = MagicMock()
        page.get_drawings.return_value = []
        page.get_image_info.return_value = []

        with patch.object(page, "find_tables") as find_tables:
            text = _extract_pymupdf_page_with_tables(page, 1, {}, "正文")

        self.assertEqual(text, "")
        find_tables.assert_not_called()

    def test_pdf_page_selection_prefers_pymupdf_when_pypdf_is_corrupted(self):
        pypdf_text = "标题 \ufffd\ufffd\ufffd\ufffd\ufffd 正文"
        pymupdf_text = "标题 完整可读的正文内容"

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pymupdf_text)

    def test_pdf_page_selection_uses_pymupdf_by_default_for_equivalent_text(self):
        pypdf_text = "The DANGER , WARNING , CAUTION , and NOTICE  statements."
        pymupdf_text = "The DANGER, WARNING, CAUTION, and NOTICE statements."

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pymupdf_text)

    def test_pdf_page_selection_prefers_more_complete_matching_text(self):
        pypdf_text = "Installation Guide\nInstall the power cable."
        pymupdf_text = (
            "Installation Guide\nInstall the power cable.\n"
            "Connect the protective earth cable before powering on the equipment."
        )

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pymupdf_text)

    def test_pdf_page_selection_keeps_pypdf_for_unrelated_extra_text(self):
        pypdf_text = "Installation Guide\nInstall the power cable."
        pymupdf_text = "Completely unrelated hidden layer " * 10

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pypdf_text)

    def test_pdf_page_selection_falls_back_when_pymupdf_is_corrupted(self):
        pypdf_text = "标题 完整可读的正文内容"
        pymupdf_text = "标题 \ufffd\ufffd\ufffd\ufffd\ufffd 正文"

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pypdf_text)

    def test_pdf_page_selection_falls_back_when_pypdf_is_more_complete(self):
        pymupdf_text = "Installation Guide\nInstall the power cable."
        pypdf_text = (
            "Installation Guide\nInstall the power cable.\n"
            "Connect the protective earth cable before powering on the equipment."
        )

        selected = _select_pdf_page_text(pypdf_text, pymupdf_text)

        self.assertEqual(selected, pypdf_text)

    def test_extracts_docx_images_with_position_based_names(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            document_path = root / "图纸.docx"
            image_dir = root / "images"
            _write_docx_with_inline_image(document_path)

            images = extract_images(
                document_path, "docx", image_dir, source_filename="图纸.docx"
            )

        self.assertEqual(len(images), 1)
        image = images[0]
        self.assertEqual(image["position"], "document-1-安装步骤-block002-p002")
        self.assertTrue(
            image["filename"].startswith("0001_document-1-安装步骤-block002-p002")
        )
        self.assertEqual(image["mime_type"], "image/png")
        self.assertIn("图纸.docx", format_image_document_text("图纸.docx", images))

    def test_image_items_from_meta_normalizes_image_list(self):
        raw = """
        {
          "images": [
            {
              "id": "image-0001",
              "filename": "0001_page001-image001.png",
              "relative_path": "task/0001_page001-image001.png",
              "mime_type": "image/png",
              "position": "page001-image001",
              "source": "报告.pdf",
              "size_bytes": 128
            }
          ]
        }
        """

        images = image_items_from_meta(raw)

        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["filename"], "0001_page001-image001.png")
        self.assertEqual(images[0]["relative_path"], "task/0001_page001-image001.png")

    def test_image_items_from_meta_tolerates_invalid_size(self):
        raw = '{"images":[{"filename":"image.png","relative_path":"task/image.png","size_bytes":"invalid"}]}'

        images = image_items_from_meta(raw)

        self.assertEqual(images[0]["size_bytes"], 0)

    def test_image_path_rejects_absolute_and_parent_paths(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)

            self.assertIsNone(
                image_path_from_item(root, {"relative_path": "/etc/passwd"})
            )
            self.assertIsNone(
                image_path_from_item(root, {"relative_path": "../outside.png"})
            )
            self.assertEqual(
                image_path_from_item(root, {"relative_path": "task/image.png"}),
                root.resolve() / "task" / "image.png",
            )

    def test_formats_page_screenshots_in_image_document_text(self):
        text = format_image_document_text(
            "报告.pdf",
            [],
            document_text="[第1页]\n表格内容",
            page_images=[
                {
                    "filename": "0001_page001-screenshot.png",
                    "position": "page001-screenshot",
                    "mime_type": "image/png",
                    "size_bytes": 2048,
                }
            ],
            page_selection={
                "total_pages": 1,
                "omitted_pages": 0,
                "strategy": "full",
            },
        )

        self.assertIn("page_screenshots: 1", text)
        self.assertIn("page_screenshot_selection", text)
        self.assertIn("0001_page001-screenshot.png", text)

    def test_select_pdf_pages_uses_candidates_and_segments_for_long_documents(self):
        selection = select_pdf_page_numbers(
            200, max_pages=10, candidate_pages=[50, 120]
        )

        self.assertEqual(selection["total_pages"], 200)
        self.assertEqual(len(selection["selected_pages"]), 10)
        self.assertIn(1, selection["selected_pages"])
        self.assertIn(50, selection["selected_pages"])
        self.assertIn(120, selection["selected_pages"])
        self.assertEqual(selection["omitted_pages"], 190)

    def test_candidate_pdf_pages_include_embedded_image_and_text_signals(self):
        pages = candidate_pdf_pages_for_image_check(
            "[第1页]\n普通正文\n\n[第2页]\n表格内容\n项目 参数 单位",
            [{"position": "page005-image001", "page_number": 5}],
        )

        self.assertEqual(pages, [2, 5])

    def test_video_helpers_format_sampling_context(self):
        self.assertTrue(allowed_video_file("安装调测.MP4"))
        self.assertFalse(allowed_video_file("安装调测.pdf"))
        self.assertEqual(video_extension_of("demo.MOV"), "mov")
        timestamps = _sample_video_timestamps(5.2, 3)

        text = format_video_document_text(
            "安装调测.mp4",
            [
                {
                    "filename": "0001_t000000000.jpg",
                    "position": "00:00.000",
                    "mime_type": "image/jpeg",
                    "size_bytes": 2048,
                },
                {
                    "filename": "0002_t000002000.jpg",
                    "position": "00:02.000",
                    "mime_type": "image/jpeg",
                    "size_bytes": 4096,
                },
            ],
            {"duration_seconds": 5.2},
        )

        self.assertEqual(len(timestamps), 3)
        self.assertEqual(timestamps[0], 0.0)
        self.assertLessEqual(timestamps[-1], 5.1)
        self.assertIn("file: 安装调测.mp4", text)
        self.assertIn("视频时长：00:05.200", text)
        self.assertIn("抽取帧数：2", text)
        self.assertIn("时间点 00:02.000", text)

    def test_video_process_output_decodes_utf8_bytes(self):
        text = _decode_process_output(b"ffmpeg: \xe2\x80\x9cinput\xe2\x80\x9d")

        self.assertIn("\u201cinput\u201d", text)

    def test_video_probe_uses_shorter_video_stream_duration(self):
        payload = {
            "streams": [{"duration": "131.057"}],
            "format": {"duration": "150.000"},
        }
        completed = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps(payload).encode("utf-8"),
            stderr=b"",
        )

        with patch(
            "app.documents.videos.subprocess.run", return_value=completed
        ) as runner:
            duration = _probe_video_duration(Path("video.mp4"))

        self.assertEqual(duration, 131.057)
        command = runner.call_args.args[0]
        self.assertIn("-select_streams", command)
        self.assertIn("v:0", command)
        self.assertIn("stream=duration:format=duration", command)

    def test_video_frame_command_hides_banner_and_selects_video_stream(self):
        completed = subprocess.CompletedProcess(
            args=[], returncode=0, stdout=b"", stderr=b""
        )

        with patch(
            "app.documents.videos.subprocess.run", return_value=completed
        ) as runner:
            _extract_frame(Path("video.mp4"), Path("frame.jpg"), 131.057)

        command = runner.call_args.args[0]
        self.assertIn("-hide_banner", command)
        self.assertEqual(command[command.index("-loglevel") + 1], "error")
        self.assertEqual(command[command.index("-map") + 1], "0:v:0")

    def test_video_frame_extraction_retries_nearby_timestamp(self):
        attempts = []

        def fake_extract_frame(video_path, destination, timestamp):
            attempts.append(timestamp)
            if timestamp == 2.55:
                raise _VideoFrameCommandError("当前采样点解码失败")
            destination.write_bytes(b"jpeg")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            with (
                patch("app.documents.videos._probe_video_duration", return_value=5.2),
                patch(
                    "app.documents.videos._extract_frame",
                    side_effect=fake_extract_frame,
                ),
            ):
                frames, selection = extract_video_frames(
                    root / "video.mp4",
                    root / "frames",
                    max_frames=3,
                )

        self.assertEqual(len(frames), 3)
        self.assertIn(2.55, attempts)
        self.assertIn(2.05, attempts)
        self.assertEqual(selection["fallback_frame_count"], 1)
        self.assertEqual(selection["skipped_frame_count"], 0)
        self.assertIn(2.05, selection["selected_timestamps"])

    def test_video_frame_extraction_uses_bounded_parallel_workers(self):
        active = 0
        maximum = 0
        lock = threading.Lock()

        def fake_extract_frame(
            _video_path, _output_dir, sequence, timestamp, _duration
        ):
            nonlocal active, maximum
            with lock:
                active += 1
                maximum = max(maximum, active)
            time.sleep(0.01)
            destination = _output_dir / f"{sequence:04d}.jpg"
            destination.write_bytes(b"jpeg")
            with lock:
                active -= 1
            return timestamp, destination.name, destination

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            with (
                patch("app.documents.videos._probe_video_duration", return_value=9.2),
                patch(
                    "app.documents.videos._extract_frame_with_fallback",
                    side_effect=fake_extract_frame,
                ),
            ):
                frames, _selection = extract_video_frames(
                    root / "video.mp4",
                    root / "frames",
                    max_frames=4,
                )

        self.assertEqual(maximum, 2)
        self.assertEqual(
            [frame["id"] for frame in frames],
            ["frame-0001", "frame-0002", "frame-0003", "frame-0004"],
        )

    def test_video_frame_extraction_skips_one_isolated_bad_sample(self):
        def fake_extract_frame(video_path, destination, timestamp):
            if 1.5 <= timestamp <= 3.6:
                raise _VideoFrameCommandError("局部视频数据损坏")
            destination.write_bytes(b"jpeg")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            with (
                patch("app.documents.videos._probe_video_duration", return_value=5.2),
                patch(
                    "app.documents.videos._extract_frame",
                    side_effect=fake_extract_frame,
                ),
            ):
                frames, selection = extract_video_frames(
                    root / "video.mp4",
                    root / "frames",
                    max_frames=3,
                )

        self.assertEqual(len(frames), 2)
        self.assertEqual(selection["skipped_frame_count"], 1)
        self.assertEqual(selection["skipped_timestamps"], [2.55])

    def test_video_frame_extraction_rejects_insufficient_coverage(self):
        def fake_extract_frame(video_path, destination, timestamp):
            if timestamp < 4.0:
                raise _VideoFrameCommandError("视频大范围无法解码")
            destination.write_bytes(b"jpeg")

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            with (
                patch("app.documents.videos._probe_video_duration", return_value=5.2),
                patch(
                    "app.documents.videos._extract_frame",
                    side_effect=fake_extract_frame,
                ),
            ):
                with self.assertRaisesRegex(
                    VideoFrameExtractionError, "仅成功抽取 1/3 帧"
                ):
                    extract_video_frames(
                        root / "video.mp4",
                        root / "frames",
                        max_frames=3,
                    )


def _write_docx_with_inline_image(path: Path):
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="xml" ContentType="application/xml"/>
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="png" ContentType="image/png"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdImage1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>
</Relationships>""",
        )
        archive.writestr(
            "word/document.xml",
            """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <w:body>
    <w:p>
      <w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
      <w:r><w:t>1 安装步骤</w:t></w:r>
    </w:p>
    <w:p>
      <w:r>
        <w:drawing>
          <a:graphic>
            <a:graphicData>
              <a:blip r:embed="rIdImage1"/>
            </a:graphicData>
          </a:graphic>
        </w:drawing>
      </w:r>
    </w:p>
  </w:body>
</w:document>""",
        )
        archive.writestr("word/media/image1.png", _TINY_PNG)


def _add_docx_hyperlink(paragraph, label: str, target: str):
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


if __name__ == "__main__":
    unittest.main()
