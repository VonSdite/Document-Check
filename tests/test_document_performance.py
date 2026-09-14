import random
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import fitz
from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.documents.extraction import extract_document, pdf, spreadsheets

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class SpreadsheetPerformanceTest(unittest.TestCase):
    def test_one_workbook_preserves_cached_shared_formulas_and_cell_types(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "types.xlsx"
            book = Workbook()
            book.epoch = CALENDAR_MAC_1904
            sheet = book.active
            sheet["A1"] = "=1+2"
            sheet["A2"] = '=HYPERLINK("https://example.test", "链接")'
            sheet["A3"] = "=SUM(A1:A2)"
            sheet["B3"] = "=SUM(B1:B2)"
            sheet["A4"] = datetime(2026, 9, 14)
            sheet["B4"] = True
            sheet["C4"] = "#DIV/0!"
            sheet["D4"] = "普通文本"
            book.save(path)
            book.close()
            with zipfile.ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            tree = ET.fromstring(members["xl/worksheets/sheet1.xml"])
            cells = {cell.get("r"): cell for cell in tree.iter(NS + "c")}
            cells["A1"].find(NS + "v").text = "3"
            cells["A2"].set("t", "str")
            cells["A2"].find(NS + "v").text = "缓存链接标题"
            cells["A3"].find(NS + "f").attrib.update(
                {"t": "shared", "si": "0", "ref": "A3:B3"}
            )
            cells["B3"].find(NS + "f").attrib.update({"t": "shared", "si": "0"})
            cells["B3"].find(NS + "f").text = None
            members["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
            with zipfile.ZipFile(path, "w") as archive:
                for name, content in members.items():
                    archive.writestr(name, content)
            with (
                patch.object(
                    spreadsheets, "load_workbook", wraps=load_workbook
                ) as load,
                patch.object(
                    ReadOnlyCell,
                    "__init__",
                    side_effect=AssertionError("解析按稀疏单元格值进行"),
                ),
            ):
                text, hyperlinks = extract_document(path, "xlsx")
            self.assertEqual(load.call_count, 1)
            self.assertIn("\n3\n", text)
            self.assertIn("缓存链接标题（超链接：https://example.test）", text)
            self.assertIn("=SUM(A1:A2) | =SUM(B1:B2)", text)
            self.assertIn("2026-09-14 00:00:00 | True | #DIV/0! | 普通文本", text)
            self.assertEqual(hyperlinks[0]["target"], "https://example.test")

    def test_indexed_hyperlinks_match_first_declared_overlapping_range(self):
        randomizer = random.Random(23)
        ranges = []
        for index in range(100):
            row, column = randomizer.randrange(1, 60), randomizer.randrange(1, 25)
            ranges.append(
                (
                    column,
                    row,
                    column + randomizer.randrange(6),
                    row + randomizer.randrange(6),
                    f"target-{index}",
                    str(index),
                )
            )
        lookup = spreadsheets._WorksheetHyperlinks(ranges)
        for row in range(1, 70):
            row_links = lookup.for_row(row)
            for column in range(1, 32):
                expected = next(
                    (
                        (target, display)
                        for left, top, right, bottom, target, display in ranges
                        if left <= column <= right and top <= row <= bottom
                    ),
                    None,
                )
                self.assertEqual(row_links.at(column), expected)

    def test_hyperlinks_on_missing_rows_and_sparse_columns_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sparse.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "开始"
            sheet["C3"] = "结束"
            sheet["A1"].hyperlink = "https://example.test/range"
            book.save(path)
            book.close()
            with zipfile.ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            tree = ET.fromstring(members["xl/worksheets/sheet1.xml"])
            tree.find(NS + "hyperlinks")[0].set("ref", "A1:B2")
            members["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
            with zipfile.ZipFile(path, "w") as archive:
                for name, content in members.items():
                    archive.writestr(name, content)
            text, links = extract_document(path, "xlsx")
            self.assertEqual(
                [link["location"].split("!")[-1] for link in links],
                ["A1", "B1", "A2", "B2"],
            )
            self.assertTrue(text.endswith(" |  | 结束"))

    def test_distant_rows_skip_blank_ranges_and_preserve_text_formula_links(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "distant.xlsx"
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "开始"
            sheet["B1000000"] = '=HYPERLINK("https://example.test", "链接")'
            sheet["B1000000"].data_type = "s"
            book.save(path)
            book.close()
            with patch.object(
                spreadsheets,
                "check_extraction_canceled",
                wraps=spreadsheets.check_extraction_canceled,
            ) as checkpoint:
                text, links = extract_document(path, "xlsx")
            self.assertLess(checkpoint.call_count, 20)
            self.assertIn("开始", text)
            self.assertEqual(links[0]["location"], "工作表“Sheet”!B1000000")
            self.assertEqual(links[0]["target"], "https://example.test")


class PdfPerformanceTest(unittest.TestCase):
    def test_blank_table_cells_skip_page_wide_textbox_scans(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sparse.pdf"
            with fitz.open() as document:
                page = document.new_page()
                for y in range(50, 251, 20):
                    page.draw_line((50, y), (450, y))
                for x in range(50, 451, 50):
                    page.draw_line((x, 50), (x, 250))
                for y in range(65, 250, 20):
                    page.insert_text((55, y), "Value", fontsize=8)
                document.save(path)
            original = fitz.Page.get_textbox
            with patch.object(
                fitz.Page, "get_textbox", autospec=True, side_effect=original
            ) as textbox:
                text, _ = extract_document(path, "pdf")
            self.assertIn("[空单元格]", text)
            self.assertIn("Value", text)
            self.assertLess(textbox.call_count, 10)

    def test_text_presence_search_is_bounded_by_nearby_characters(self):
        chars = [{"bbox": (0, index * 10, 5, index * 10 + 5)} for index in range(10000)]
        index = pdf._PdfTextPresenceIndex(
            {"blocks": [{"lines": [{"spans": [{"chars": chars}]}]}]}
        )
        with patch.object(
            pdf, "_pdf_rects_intersect", wraps=pdf._pdf_rects_intersect
        ) as intersects:
            self.assertTrue(index.may_have_text((0, 50000, 5, 50005)))
            self.assertFalse(index.may_have_text((50, 50000, 55, 50005)))
        self.assertLessEqual(intersects.call_count, 4)
        self.assertTrue(pdf._PdfTextPresenceIndex({}).may_have_text((0, 0, 5, 5)))
        incomplete = {"blocks": [{"lines": [{"spans": [{"text": "候选"}]}]}]}
        self.assertTrue(
            pdf._PdfTextPresenceIndex(incomplete).may_have_text((0, 0, 5, 5))
        )

    def test_nontext_index_preserves_image_and_drawing_boundaries(self):
        randomizer = random.Random(37)
        images = []
        drawings = []
        for _ in range(200):
            x, y = randomizer.uniform(0, 100), randomizer.uniform(0, 100)
            rect = (x, y, x + randomizer.uniform(0, 10), y + randomizer.uniform(0, 10))
            images.append(rect)
            drawings.append({"rect": rect})
        index = pdf._PdfNontextIndex(images, drawings)
        for _ in range(200):
            x, y = randomizer.uniform(0, 100), randomizer.uniform(0, 100)
            bbox = (x, y, x + 12, y + 8)
            interior = (x + 1, y + 1, x + 11, y + 7)
            expected = any(
                (rect[2] - rect[0]) * (rect[3] - rect[1]) >= 4
                and pdf._pdf_rect_contains_center(bbox, rect)
                for rect in images
            ) or any(
                rect[2] - rect[0] >= 2
                and rect[3] - rect[1] >= 2
                and pdf._pdf_rect_contains_center(interior, rect)
                for rect in (drawing["rect"] for drawing in drawings)
            )
            self.assertEqual(index.contains(bbox), expected)


if __name__ == "__main__":
    unittest.main()
