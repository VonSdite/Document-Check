import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from app.checks.agent_policy import POLICY_KEY, check_agent_options, strategy_for
from app.contracts.agent import AgentOptions
from app.documents.index import DocumentIndex
from app.models.agent import AgentTurn
from app.models.client import LLMError
from app.persistence.checkpoints import (
    CheckpointOwnershipLost,
    CheckpointStore,
    MemoryCheckpointStore,
)
from app.persistence.connection import get_db
from app.persistence.defaults import seed_defaults
from app.persistence.schema import init_db
from app.persistence.settings import set_setting
from app.reporting.excel import build_report_workbook
from app.reporting.service import (
    _matching_report_suppression_rule,
    _prepare_task_results,
)
from app.tasks.runner import TaskRunner
from app.tasks.runtime.evidence_agent import EvidenceAgent, workflow_signature
from app.tasks.selection import _check_items_from_snapshot
from app.tasks.submission import _enabled_check_item_snapshots
from tests.test_evidence_agent import (
    ScriptedModel,
    discovery,
    empty_plan,
    input_context,
)


class EvidenceIntegrationTest(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.TemporaryDirectory()
        self.app = Flask(__name__)
        self.app.config.update(
            DATABASE=str(Path(self.root.name) / "test.sqlite3"),
            UPLOAD_FOLDER=str(Path(self.root.name) / "uploads"),
            IMAGE_FOLDER=str(Path(self.root.name) / "images"),
            NETWORK={"proxy_mode": "direct", "proxy": "", "ssl_verify": False},
        )
        Path(self.app.config["UPLOAD_FOLDER"]).mkdir()
        Path(self.app.config["IMAGE_FOLDER"]).mkdir()
        self.context = self.app.app_context()
        self.context.push()
        init_db()
        seed_defaults()

    def tearDown(self):
        self.context.pop()
        self.root.cleanup()

    def insert_task(self, snapshots, *, token="owner", text=None):
        db = get_db()
        cursor = db.execute(
            """INSERT INTO tasks(ip, original_filename, stored_filename, file_type, file_size,
            checks_json, checks_snapshot_json, model_name, api_base, status, claim_token, document_text, created_at, updated_at)
            VALUES ('127.0.0.1', 'a.txt', 'a.txt', 'txt', 10, '[]', ?, 'test', 'http://model/v1/chat/completions', 'running', ?, ?, '2026-01-01', '2026-01-01')""",
            (json.dumps(snapshots, ensure_ascii=False), token, text),
        )
        db.commit()
        return cursor.lastrowid

    def test_text_task_policies_use_built_in_defaults_and_ignore_public_overrides(
        self,
    ):
        for task_type in (
            "document_check",
            "consistency_check",
            "language_consistency_check",
        ):
            row = (
                get_db()
                .execute(
                    "SELECT * FROM check_items WHERE task_type = ? AND enabled=1 LIMIT 1",
                    (task_type,),
                )
                .fetchone()
            )
            if row is None:
                self.fail("默认检查项缺失")
            public_override = AgentOptions(
                strategy="compliance",
                max_calls=10,
                standards=[{"name": "规范", "version": "1", "text": "电压必须明确。"}],
            ).model_dump(mode="json")
            set_setting(POLICY_KEY + row["code"], public_override)
            snapshot = _enabled_check_item_snapshots(get_db(), [row["id"]], task_type)
            self.assertEqual(
                _check_items_from_snapshot(json.dumps(snapshot))[0]["agent_options"],
                AgentOptions().model_dump(mode="json"),
            )

    def test_checkpoint_write_requires_current_execution_owner(self):
        task_id = self.insert_task([])
        store = CheckpointStore(task_id, "check", "signature", "owner")
        store.save({"next": "discover"})
        self.assertEqual(store.load()["next"], "discover")
        get_db().execute(
            "UPDATE tasks SET claim_token='new-owner' WHERE id=?", (task_id,)
        )
        get_db().commit()
        with self.assertRaises(CheckpointOwnershipLost):
            store.save({"next": "done"})
        self.assertEqual(store.load()["next"], "discover")

    def test_checkpoint_compatibility_keeps_task_item_and_execution_owner_boundaries(
        self,
    ):
        task_id = self.insert_task([])
        old_store = CheckpointStore(task_id, "check", "previous", "owner")
        old_store.save({"next": "discover"})
        store = CheckpointStore(
            task_id, "check", "current", "owner", compatible_signatures=("previous",)
        )
        self.assertEqual(store.load(), {"next": "discover"})
        self.assertIsNone(
            CheckpointStore(task_id, "check", "different", "owner").load()
        )
        self.assertIsNone(
            CheckpointStore(
                task_id,
                "other-check",
                "current",
                "owner",
                compatible_signatures=("previous",),
            ).load()
        )
        other_task = self.insert_task([])
        self.assertIsNone(
            CheckpointStore(
                other_task,
                "check",
                "current",
                "owner",
                compatible_signatures=("previous",),
            ).load()
        )
        store.save({"next": "review"})
        self.assertIsNone(old_store.load())
        self.assertEqual(store.load(), {"next": "review"})
        for assignment in (
            "cancel_requested=1",
            "cancel_requested=0, claim_token='other-owner'",
        ):
            get_db().execute(f"UPDATE tasks SET {assignment} WHERE id=?", (task_id,))
            get_db().commit()
            with self.assertRaises(CheckpointOwnershipLost):
                store.save({"next": "done"})
            self.assertEqual(store.load(), {"next": "review"})

    def test_worker_resumes_legacy_checkpoint_beyond_elapsed_and_call_limits(self):
        text = "\n".join(f"来源{n}：" + "设备参数。" * 120 for n in range(3))
        meta = {
            "source_documents": [
                {"filename": "a.txt", "text": text, "text_complete": True}
            ]
        }
        index = DocumentIndex.from_task(text, meta, "a.txt")
        old_options = AgentOptions(
            protocol="structured", chunk_chars=1000, max_calls=96, max_seconds=3600
        )
        for missing_limits in (False, True):
            with self.subTest(missing_limits=missing_limits):
                snapshot_options = old_options.model_dump(mode="json")
                if missing_limits:
                    snapshot_options.pop("max_calls")
                    snapshot_options.pop("max_seconds")
                item = {
                    "code": "custom-check",
                    "name": "业务检查",
                    "prompt": "检查设备参数。",
                    "agent_options": snapshot_options,
                }
                task_id = self.insert_task([item], text=text)
                get_db().execute(
                    "UPDATE tasks SET document_meta_json=? WHERE id=?",
                    (json.dumps(meta, ensure_ascii=False), task_id),
                )
                get_db().commit()
                connection = dict(
                    get_db()
                    .execute("SELECT * FROM tasks WHERE id=?", (task_id,))
                    .fetchone()
                )
                legacy_signature = workflow_signature(
                    index,
                    item,
                    connection,
                    old_options,
                    "document_check",
                    include_execution_limits=True,
                )
                memory = MemoryCheckpointStore()
                discoveries = 0

                def handler(messages, title, tools, force):
                    nonlocal discoveries
                    if title == "CheckPlan":
                        return empty_plan()
                    discoveries += 1
                    return (
                        discovery(input_context(messages))
                        if discoveries == 1
                        else LLMError("503 服务暂时失败")
                    )

                first = EvidenceAgent(
                    index=index,
                    item=item,
                    task_type="document_check",
                    connection=connection,
                    options=old_options,
                    parallelism=1,
                    store=memory,
                    ensure_active=lambda: None,
                    model_factory=lambda connection, **kwargs: ScriptedModel(
                        handler, connection, **kwargs
                    ),
                ).run()
                self.assertIn("503", first["error"])
                state = memory.load()
                self.assertEqual(len(state["checked_ids"]), 1)
                state["elapsed_seconds"] = 3601.0
                state["requests"].extend(dict(state["requests"][0]) for _ in range(93))
                CheckpointStore(task_id, item["code"], legacy_signature, "owner").save(
                    state
                )
                known_input = sum(r["input_tokens"] or 0 for r in state["requests"])
                calls = []
                checked = []

                def turn(model, messages, *, schema, **kwargs):
                    self.assertEqual(model.connection["task_id"], task_id)
                    self.assertEqual(schema["title"], "Discovery")
                    calls.append(schema["title"])
                    model.before_request({"messages": messages})
                    self.assertEqual(
                        model.connection["remaining_seconds"],
                        connection["request_timeout"],
                    )
                    context = input_context(messages)
                    checked.extend(s["id"] for s in context["sources"])
                    response = AgentTurn(
                        json.dumps(discovery(context), ensure_ascii=False),
                        usage={"prompt_tokens": 100, "completion_tokens": 50},
                        finish_reason="stop",
                    )
                    model.after_request({}, response, 0.01, None)
                    return response

                with patch("app.models.agent.AgentModel.turn", new=turn):
                    TaskRunner(self.app).run(task_id, "owner")
                task = (
                    get_db()
                    .execute("SELECT * FROM tasks WHERE id=?", (task_id,))
                    .fetchone()
                )
                self.assertEqual(task["status"], "completed", task["error"])
                result = json.loads(task["result_json"])[0]
                self.assertTrue(result["agent"]["coverage"]["all_checked"])
                self.assertEqual(len(calls), 2)
                self.assertEqual(len(set(checked)), 2)
                self.assertFalse(set(checked) & set(state["checked_ids"]))
                self.assertEqual(result["agent"]["model_calls"], 98)
                self.assertEqual(
                    result["agent"]["known_input_tokens"], known_input + 200
                )
                self.assertGreaterEqual(result["agent"]["elapsed_seconds"], 3601.0)
                current_signature = workflow_signature(
                    index,
                    item,
                    connection,
                    old_options.model_copy(
                        update={"max_calls": None, "max_seconds": None}
                    ),
                    "document_check",
                )
                self.assertNotEqual(legacy_signature, current_signature)
                self.assertEqual(
                    CheckpointStore(
                        task_id, item["code"], current_signature, "owner"
                    ).load()["next"],
                    "done",
                )
                self.assertIsNone(
                    CheckpointStore(
                        task_id, item["code"], legacy_signature, "owner"
                    ).load()
                )

    def test_end_to_end_worker_indexes_long_document_and_saves_valid_report(self):
        text = "\n".join("设备参数正常。" * 500 for _ in range(28))
        (Path(self.app.config["UPLOAD_FOLDER"]) / "a.txt").write_text(
            text, encoding="utf-8"
        )
        snapshots = [
            {
                "code": "custom-check",
                "name": "业务检查",
                "prompt": "检查设备参数。",
                "agent_options": {
                    **AgentOptions(protocol="structured").model_dump(mode="json"),
                    "max_output_tokens": 2048,
                    "max_tokens": 500000,
                },
            }
        ]
        get_db().execute("ALTER TABLE tasks ADD COLUMN max_output_tokens INTEGER")
        task_id = self.insert_task(snapshots)
        get_db().execute(
            "UPDATE tasks SET max_output_tokens=65536 WHERE id=?", (task_id,)
        )
        get_db().commit()
        calls = []

        def turn(model, messages, *, schema, **kwargs):
            calls.append(kwargs)
            model.before_request({"messages": messages})
            answer = (
                empty_plan()
                if schema["title"] == "CheckPlan"
                else discovery(input_context(messages))
            )
            response = AgentTurn(
                json.dumps(answer, ensure_ascii=False),
                usage={"prompt_tokens": 350000, "completion_tokens": 300000},
                finish_reason="stop",
            )
            model.after_request({}, response, 0.01, None)
            return response

        with patch("app.models.agent.AgentModel.turn", new=turn):
            TaskRunner(self.app).run(task_id, "owner")
        task = get_db().execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        self.assertEqual(task["status"], "completed", task["error"])
        self.assertTrue(calls)
        self.assertTrue(all("max_output_tokens" not in call for call in calls))
        result = json.loads(task["result_json"])[0]
        self.assertEqual(result["structured_report"]["items"], [])
        self.assertEqual(result["agent"]["coverage"]["checked_blocks"], 28)
        self.assertEqual(result["agent"]["known_input_tokens"], len(calls) * 350000)
        self.assertEqual(result["agent"]["known_output_tokens"], len(calls) * 300000)
        self.assertIsNotNone(
            get_db()
            .execute("SELECT 1 FROM task_checkpoints WHERE task_id=?", (task_id,))
            .fetchone()
        )
        meta = json.loads(task["document_meta_json"])
        self.assertTrue(meta["source_documents"][0]["text_complete"])
        output = build_report_workbook(task)
        self.assertGreater(len(output.getvalue()), 1000)
        workbook = load_workbook(output)
        self.assertEqual(workbook["执行用量"].cell(2, 3).value, len(calls) * 350000)
        self.assertEqual(workbook["执行用量"].cell(2, 4).value, len(calls) * 300000)

    def test_invalid_report_is_not_recovered_as_unvalidated_text(self):
        result = {
            "code": "check",
            "name": "检查",
            "agent": {"workflow_version": "v1"},
            "structured_report": {"items": "invalid"},
            "result": "1. 错误：凭空内容",
        }
        prepared = _prepare_task_results([result], suppression_rules={})[0]
        self.assertEqual(prepared["report_items"], [])
        self.assertIn("校验失败", prepared["error"])

    def test_feedback_suppression_requires_matching_evidence_context(self):
        item = {"description": "电压不一致", "evidence_fingerprint": "context-a"}
        rule = {
            "id": 1,
            "description": "电压不一致",
            "evidence_fingerprint": "context-b",
        }
        self.assertIsNone(_matching_report_suppression_rule(item, [rule])[0])
        rule["evidence_fingerprint"] = "context-a"
        self.assertEqual(_matching_report_suppression_rule(item, [rule])[0]["id"], 1)
        rule.pop("evidence_fingerprint")
        self.assertIsNone(_matching_report_suppression_rule(item, [rule])[0])

    def test_custom_checks_plan_scope_with_built_in_policies(self):
        options = check_agent_options("custom-business-check", "document_check")
        self.assertEqual(options, AgentOptions().model_dump(mode="json"))
        self.assertEqual(
            strategy_for("custom-business-check", "document_check", "auto"), "auto"
        )
        self.assertEqual(strategy_for("typo", "document_check", "auto"), "local")
        self.assertEqual(
            strategy_for("custom", "consistency_check", "auto"), "cross_document"
        )
        self.assertEqual(
            strategy_for("custom", "language_consistency_check", "auto"),
            "cross_language",
        )
        self.assertIsNone(check_agent_options("common-terms", "document_check"))


if __name__ == "__main__":
    unittest.main()
