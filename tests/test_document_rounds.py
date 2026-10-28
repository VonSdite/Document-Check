import json
import tempfile
import unittest
from pathlib import Path
from threading import Barrier, Lock
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from app.checks.common_terms import COMMON_TERMS_CHECK_CODE
from app.checks.hyperlinks import HYPERLINK_CHECK_CODE
from app.checks.sensitive_terms import SENSITIVE_TERMS_CHECK_CODE
from app.models.client import LLMError
from app.persistence.connection import get_db, now_text
from app.persistence.schema import init_db
from app.persistence.settings import set_setting
from app.reporting.excel import build_report_workbook
from app.reporting.rounds import build_round_result
from app.reporting.service import _prepare_task_results
from app.tasks.activity import request_check_cancellation
from app.tasks.runner import TaskRunner


def finding(description="电压单位不一致", **overrides):
    return {
        "status": "issue",
        "severity": "medium",
        "confidence": "high",
        "category": "单位",
        "location": "第 2 节",
        "excerpt": "输入电压 5mV",
        "description": description,
        "impact": "影响操作准确性。",
        "suggestion": "核对并统一单位。",
        **overrides,
    }


def report(*items):
    return json.dumps(
        {"summary": "检查完成。", "items": list(items)}, ensure_ascii=False
    )


def completed(number, *items):
    return {
        "round": number,
        "status": "completed",
        "result": report(*items),
        "report": json.loads(report(*items)),
    }


class RoundAggregationTest(unittest.TestCase):
    def test_merge_paraphrases_keep_single_round_findings_and_count_distinct_rounds(
        self,
    ):
        result = build_round_result(
            "check",
            3,
            [
                completed(1, finding(), finding()),
                completed(2, finding("文中电压单位不一致")),
                completed(
                    3,
                    finding(
                        "操作步骤缺少退出说明",
                        category="步骤",
                        excerpt="完成配置后退出",
                        location="第 3 节",
                    ),
                ),
            ],
        )
        items = result["structured_report"]["items"]
        self.assertEqual(len(items), 2)
        self.assertEqual(items[0]["source_rounds"], [1, 2])
        self.assertEqual(items[1]["source_rounds"], [3])
        self.assertEqual(items[0]["confidence"], "high")
        self.assertEqual(len({item["id"] for item in items}), 2)

    def test_same_quote_with_different_issues_and_numbers_stays_separate(self):
        result = build_round_result(
            "check",
            2,
            [
                completed(1, finding("电压应为5V"), finding("参数标点格式错误")),
                completed(2, finding("电压应为12V")),
            ],
        )
        self.assertEqual(len(result["structured_report"]["items"]), 3)

    def test_conflicting_judgments_require_confirmation(self):
        result = build_round_result(
            "check",
            2,
            [
                completed(1, finding()),
                completed(2, finding(status="suggestion")),
            ],
        )
        item = result["structured_report"]["items"][0]
        self.assertTrue(item["round_conflict"])
        self.assertEqual(item["status"], "suggestion")
        self.assertEqual(item["confidence"], "low")

    def test_conflicting_numeric_fixes_require_confirmation(self):
        result = build_round_result(
            "check",
            2,
            [
                completed(1, finding(suggestion="将电压改为5V。")),
                completed(2, finding(suggestion="将电压改为12V。", severity="high")),
            ],
        )
        item = result["structured_report"]["items"][0]
        self.assertTrue(item["round_conflict"])
        self.assertEqual(item["status"], "suggestion")
        self.assertEqual(item["severity"], "high")
        self.assertIn("5V", item["suggestion"])
        self.assertIn("12V", item["suggestion"])

    def test_merged_report_preserves_more_than_thirty_findings_and_sources(self):
        result = build_round_result(
            "check",
            2,
            [
                completed(
                    number,
                    *[
                        finding(
                            f"问题 {number}-{index}", excerpt=f"原文 {number}-{index}"
                        )
                        for index in range(25)
                    ],
                )
                for number in (1, 2)
            ],
        )
        prepared = _prepare_task_results(
            [{"code": "check", "name": "检查", "issue_output_limit": 30, **result}],
            suppression_rules={},
        )[0]
        self.assertEqual(len(prepared["report_items"]), 50)
        self.assertIsNone(prepared["report_limit"])
        self.assertEqual(prepared["report_items"][0]["source_rounds"], [1])


class DocumentRoundExecutionTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.app = Flask(__name__, instance_path=str(root / "instance"))
        self.app.config.update(
            DATABASE=str(root / "test.sqlite3"),
            UPLOAD_FOLDER=str(root / "uploads"),
            IMAGE_FOLDER=str(root / "images"),
            NETWORK={"proxy_mode": "direct", "proxy": "", "ssl_verify": False},
        )
        (root / "uploads").mkdir()
        self.context = self.app.app_context()
        self.context.push()
        init_db()

    def tearDown(self):
        self.context.pop()
        self.temp_dir.cleanup()

    def create_task(self, codes=("compliance",), rounds=5):
        checks = [
            {"id": index, "code": code, "name": code, "prompt": "检查全文"}
            for index, code in enumerate(codes, 1)
        ]
        cursor = get_db().execute(
            "INSERT INTO tasks(ip, owner_subject, original_filename, stored_filename, "
            "file_type, file_size, document_text, document_meta_json, checks_json, "
            "checks_snapshot_json, model_name, api_base, api_key, status, created_at, updated_at) "
            "VALUES ('127.0.0.1', 'ip:127.0.0.1', 'doc.txt', 'doc.txt', 'txt', 10, ?, ?, ?, ?, "
            "'test-model', 'https://example.test', 'secret', 'running', ?, ?)",
            (
                "输入电压 5mV\n完整文档正文",
                json.dumps({"check_rounds": rounds}),
                json.dumps([item["id"] for item in checks]),
                json.dumps(checks),
                now_text(),
                now_text(),
            ),
        )
        get_db().commit()
        return cursor.lastrowid

    def task(self, task_id):
        return dict(
            get_db().execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        )

    def retry(self, task_id):
        get_db().execute(
            "UPDATE tasks SET status='running', cancel_requested=0, api_key='secret', "
            "retry_check_codes_json='[\"compliance\"]' WHERE id=?",
            (task_id,),
        )
        get_db().commit()

    def test_five_independent_requests_reuse_text_and_keep_each_raw_reply(self):
        task_id = self.create_task()
        with patch(
            "app.tasks.runner.run_check", return_value=report(finding())
        ) as model:
            TaskRunner(self.app).run(task_id)
        task = self.task(task_id)
        self.assertEqual(task["status"], "completed")
        self.assertIsNone(task["api_key"])
        self.assertEqual(model.call_count, 5)
        for call in model.call_args_list:
            self.assertEqual(call.kwargs["document_text"], task["document_text"])
            self.assertEqual(call.kwargs["prompt"], "检查全文")
        result = json.loads(task["result_json"])[0]
        self.assertEqual(
            result["round_summary"], {"planned": 5, "completed": 5, "failed": 0}
        )
        self.assertEqual(len(result["round_results"]), 5)
        self.assertEqual(
            result["structured_report"]["items"][0]["source_rounds"], [1, 2, 3, 4, 5]
        )

    def test_partial_failure_retry_runs_only_missing_round_and_keeps_review_ids(self):
        task_id = self.create_task()
        with patch(
            "app.tasks.runner.run_check",
            side_effect=[
                report(finding()),
                LLMError("暂时不可用"),
                report(),
                report(),
                report(),
            ],
        ):
            TaskRunner(self.app).run(task_id)
        first = self.task(task_id)
        self.assertEqual(first["status"], "partial")
        result = json.loads(first["result_json"])[0]
        item_id = result["structured_report"]["items"][0]["id"]
        result["item_acceptances"] = {item_id: {"status": "accepted"}}
        get_db().execute(
            "UPDATE tasks SET result_json=? WHERE id=?", (json.dumps([result]), task_id)
        )
        get_db().commit()
        self.retry(task_id)
        set_setting("document_check_rounds", 1)
        with patch(
            "app.tasks.runner.run_check", return_value=report(finding())
        ) as model:
            TaskRunner(self.app).run(task_id)
        final = self.task(task_id)
        self.assertEqual(final["status"], "completed")
        model.assert_called_once()
        result = json.loads(final["result_json"])[0]
        self.assertEqual(result["round_summary"]["completed"], 5)
        self.assertEqual(result["structured_report"]["items"][0]["id"], item_id)
        self.assertIn(item_id, result["item_acceptances"])
        self.assertEqual(
            result["structured_report"]["items"][0]["source_rounds"], [1, 2]
        )

    def test_invalid_or_truncated_round_is_failure_and_other_rounds_continue(self):
        task_id = self.create_task(rounds=3)
        with patch(
            "app.tasks.runner.run_check",
            side_effect=[
                '{"items":[{"description":"未完成"}',
                '{"summary":"空报告"}',
                report(),
            ],
        ):
            TaskRunner(self.app).run(task_id)
        task = self.task(task_id)
        self.assertEqual(task["status"], "partial")
        self.assertEqual(
            json.loads(task["result_json"])[0]["round_summary"],
            {"planned": 3, "completed": 1, "failed": 2},
        )

    def test_retrying_first_round_keeps_ids_from_later_successful_rounds(self):
        task_id = self.create_task(rounds=2)
        with patch(
            "app.tasks.runner.run_check",
            side_effect=[LLMError("失败"), report(finding())],
        ):
            TaskRunner(self.app).run(task_id)
        original = json.loads(self.task(task_id)["result_json"])[0]
        original_id = original["structured_report"]["items"][0]["id"]
        self.retry(task_id)
        with patch(
            "app.tasks.runner.run_check",
            return_value=report(finding("文中电压单位不一致")),
        ):
            TaskRunner(self.app).run(task_id)
        final = json.loads(self.task(task_id)["result_json"])[0]
        self.assertEqual(len(final["structured_report"]["items"]), 1)
        self.assertEqual(final["structured_report"]["items"][0]["id"], original_id)

    def test_whole_task_cancellation_keeps_completed_rounds_available_for_retry(self):
        from app.tasks.retries import retry_check_codes_for_task

        task_id = self.create_task(rounds=3)
        calls = 0

        def model(**kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                get_db().execute(
                    "UPDATE tasks SET cancel_requested=1, status='canceling' WHERE id=?",
                    (task_id,),
                )
                get_db().commit()
                kwargs["check_canceled"]()
            return report(finding())

        with patch("app.tasks.runner.run_check", side_effect=model):
            TaskRunner(self.app).run(task_id)
        task = self.task(task_id)
        self.assertEqual(task["status"], "canceled")
        result = json.loads(task["result_json"])[0]
        self.assertEqual(result["round_summary"]["completed"], 1)
        self.assertEqual(retry_check_codes_for_task(task), ["compliance"])
        self.retry(task_id)
        with patch("app.tasks.runner.run_check", return_value=report()) as model:
            TaskRunner(self.app).run(task_id)
        self.assertEqual(model.call_count, 2)
        self.assertEqual(self.task(task_id)["status"], "completed")

    def test_all_failed_rounds_mark_task_failed(self):
        task_id = self.create_task(rounds=2)
        with patch("app.tasks.runner.run_check", side_effect=LLMError("失败")):
            TaskRunner(self.app).run(task_id)
        self.assertEqual(self.task(task_id)["status"], "failed")

    def test_cancel_preserves_successful_round_and_retry_fills_remaining_rounds(self):
        task_id = self.create_task(rounds=3)
        calls = 0

        def model(**kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                request_check_cancellation(task_id, None, "compliance")
                kwargs["cancel_event"].set()
                kwargs["check_canceled"]()
            return report(finding())

        with patch("app.tasks.runner.run_check", side_effect=model):
            TaskRunner(self.app).run(task_id)
        task = self.task(task_id)
        self.assertEqual(task["status"], "canceled")
        self.assertEqual(calls, 2)
        self.assertEqual(
            json.loads(task["result_json"])[0]["round_summary"]["completed"], 1
        )
        self.retry(task_id)
        with patch("app.tasks.runner.run_check", return_value=report()) as model:
            TaskRunner(self.app).run(task_id)
        self.assertEqual(model.call_count, 2)
        self.assertEqual(self.task(task_id)["status"], "completed")

    def test_rounds_share_existing_check_concurrency_limit(self):
        task_id = self.create_task(codes=("first", "second"), rounds=3)
        set_setting("check_item_concurrency", 2)
        barrier, lock = Barrier(2), Lock()
        active = maximum = 0

        def model(**kwargs):
            nonlocal active, maximum
            with lock:
                active += 1
                maximum = max(maximum, active)
            barrier.wait(timeout=5)
            with lock:
                active -= 1
            return report()

        with patch("app.tasks.runner.run_check", side_effect=model) as mocked:
            TaskRunner(self.app).run(task_id)
        self.assertEqual(mocked.call_count, 6)
        self.assertEqual(maximum, 2)
        self.assertEqual(self.task(task_id)["status"], "completed")

    def test_local_rules_run_once_while_ai_checks_repeat(self):
        task_id = self.create_task(
            codes=(
                "compliance",
                SENSITIVE_TERMS_CHECK_CODE,
                COMMON_TERMS_CHECK_CODE,
                HYPERLINK_CHECK_CODE,
            )
        )
        empty = {"summary": "检查完成", "items": []}
        with (
            patch("app.tasks.runner.run_check", return_value=report()) as model,
            patch(
                "app.tasks.runner._run_sensitive_terms_check", return_value=empty
            ) as sensitive,
            patch(
                "app.tasks.runner._run_common_terms_check", return_value=empty
            ) as common,
            patch(
                "app.tasks.runner.build_hyperlink_report", return_value=empty
            ) as links,
        ):
            TaskRunner(self.app).run(task_id)
        self.assertEqual(model.call_count, 5)
        sensitive.assert_called_once()
        common.assert_called_once()
        links.assert_called_once()
        self.assertEqual(self.task(task_id)["status"], "completed")

    def test_excel_exports_all_findings_sources_and_round_results(self):
        task_id = self.create_task(rounds=2)
        outputs = [
            report(
                *[
                    finding(f"问题 {number}-{i}", excerpt=f"证据 {number}-{i}")
                    for i in range(20)
                ]
            )
            for number in (1, 2)
        ]
        with patch("app.tasks.runner.run_check", side_effect=outputs):
            TaskRunner(self.app).run(task_id)
        workbook = load_workbook(build_report_workbook(self.task(task_id)))
        self.addCleanup(workbook.close)
        sheet = workbook["报告条目"]
        self.assertEqual(sheet.max_row, 41)
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("检查轮次", headers)
        self.assertIn("发现轮次", headers)
        self.assertEqual(sheet.cell(2, headers.index("发现轮次") + 1).value, "1")
        self.assertEqual(workbook["检查轮次"].max_row, 3)


if __name__ == "__main__":
    unittest.main()
