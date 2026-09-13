import re

from openpyxl.styles import Font, PatternFill

REPORT_ITEM_TYPES = {
    "issue": "问题",
    "suggestion": "建议",
    "non_issue": "非问题",
}
REPORT_ITEM_TYPE_ORDER = ("issue", "suggestion", "non_issue")
REPORT_COUNT_KEYS = REPORT_ITEM_TYPE_ORDER + (
    "accepted_issue",
    "rejected_issue",
    "pending_issue_acceptance",
    "suppressed",
    "reviewed",
    "pending_review",
)
REPORT_REVIEW_STATUSES = {
    "pending": "未标注",
    "in_progress": "标注中",
    "completed": "已标注",
    "empty": "无可标注条目",
    "unavailable": "—",
}
REPORT_REVIEW_FILTERS = {
    "pending": "待标注",
    "in_progress": "标注中",
    "completed": "已标注",
    "empty": "无可标注条目",
}
REPORT_ACCEPTANCE_STATUSES = {
    "pending": "未确认",
    "accepted": "接纳",
    "rejected": "不接纳",
}
REPORT_REJECTION_REASONS = {
    "false_positive": "模型误报",
    "model_hallucination": "模型幻觉",
    "evidence_insufficient": "证据不足",
    "not_applicable": "不适用",
    "other": "其他",
}
REPORT_REJECTION_REASON_HINTS = {
    "false_positive": "能找到模型依据，但结论判断不成立",
    "model_hallucination": "找不到模型引用的原文、位置或事实",
    "evidence_insufficient": "信息不完整，暂时无法判断结论对错",
    "not_applicable": "这条检查规则不适用于当前文档或场景",
    "other": "无法归入以上原因，请补充说明",
}
REPORT_SUPPRESSION_REJECTION_REASONS = {
    "model_hallucination",
    "false_positive",
    "not_applicable",
}
REPORT_SUPPRESSION_DESCRIPTION_SIMILARITY_THRESHOLD = 0.56
REPORT_STATS_PREPARATION_VERSION = "3"
# 报告统计按批次在后台刷新，请求线程只同步处理少量记录。
REPORT_STATS_INLINE_REBUILD_LIMIT = 20
REPORT_STATS_BACKGROUND_BATCH_SIZE = 100
REPORT_SUPPRESSION_DESCRIPTION_REPLACEMENTS = (
    ("不统一", "不一致"),
    ("不相同", "不一致"),
    ("存在差异", "不一致"),
    ("矛盾", "冲突"),
    ("有误", "错误"),
    ("不正确", "错误"),
    ("没有提供", "缺失"),
    ("未提供", "缺失"),
    ("没有说明", "缺失"),
    ("未说明", "缺失"),
    ("缺少", "缺失"),
    ("遗留", "保留"),
    ("残留", "保留"),
    ("资料", "文档"),
)
REPORT_ITEM_FIELDS = (
    ("severity_label", "严重程度"),
    ("confidence_label", "证据可信度"),
    ("category", "问题类型"),
    ("location", "位置"),
    ("excerpt", "原文/证据"),
    ("description", "问题描述"),
    ("impact", "影响"),
    ("suggestion", "修改建议"),
)
REPORT_SUPPRESSION_FIELDS = (
    "category",
    "location",
    "excerpt",
    "description",
    "impact",
    "suggestion",
)
REPORT_SEVERITY_LABELS = {
    "critical": "致命",
    "high": "高",
    "medium": "中",
    "low": "低",
}
REPORT_CONFIDENCE_LABELS = {
    "high": "高",
    "medium": "中",
    "low": "低",
}
REPORT_SEVERITY_ORDER = {
    value: index for index, value in enumerate(REPORT_SEVERITY_LABELS)
}
REPORT_CONFIDENCE_ORDER = {
    value: index for index, value in enumerate(REPORT_CONFIDENCE_LABELS)
}
MEDIA_REPORT_ITEM_FIELDS = (("media_summary", "AI检查结论"),)
MEDIA_REPORT_ITEM_DETAIL_FIELDS = (
    ("category", "问题类型"),
    ("location", "位置/画面"),
    ("excerpt", "依据/证据"),
    ("impact", "影响"),
    ("suggestion", "建议"),
)
REPORT_EXPORT_MIMETYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
REPORT_EXPORT_SHEET_NAME = "报告条目"
REPORT_EXPORT_RESULT_CODE_HEADER = "检查项编码（请勿修改）"
REPORT_EXPORT_ITEM_ID_HEADER = "条目标识（请勿修改）"
REPORT_IMPORT_MAX_BYTES = 10 * 1024 * 1024
REPORT_IMPORT_MAX_ROWS = 5000
REPORT_EXPORT_HEADER_FILL = PatternFill("solid", fgColor="EAF0F8")
REPORT_EXPORT_HEADER_FONT = Font(bold=True)
REPORT_EXPORT_EDITABLE_FILL = PatternFill("solid", fgColor="FFF4CC")
REPORT_TOTAL_EXPORT_ROWS = (
    ("问题", "issue"),
    ("建议", "suggestion"),
    ("非问题", "non_issue"),
    ("接纳问题", "accepted_issue"),
    ("不接纳问题", "rejected_issue"),
    ("待确认问题", "pending_issue_acceptance"),
    ("已忽略误报", "suppressed"),
    ("问题检出率", "issue_detection_rate"),
    ("问题接纳率", "issue_acceptance_rate"),
    ("合计", "total"),
)


class ReportExcelImportError(ValueError):
    pass


REPORT_ITEM_TYPE_LABEL = "条目判定"
REPORT_ITEM_START_RE = re.compile(
    r"^(?:(?:问题|建议|风险|疑点|不一致|偏差|错误|缺失)\s*\d*[:：]|"
    r"(?:\d{1,3}[.、)]|\(\d{1,3}\)|（\d{1,3}）)\s*(?:[*_`~]{1,3}\s*)?\S)"
)
REPORT_ITEM_PREFIX_RE = re.compile(
    r"^\s*(?:(?:[-+*]\s+)|(?:#{1,6}\s+)|(?:[*_`~]{1,3}\s*))*"
)
REPORT_JSON_ITEM_KEYS = ("items", "issues", "report_items", "findings", "problems")
REPORT_JSON_SUMMARY_KEYS = ("summary", "overall", "conclusion", "总体结论", "总结")
REPORT_STATUS_KEYS = (
    "status",
    "状态",
    "classification",
    "item_type",
    "结论类型",
    "问题状态",
    "type",
)
REPORT_FIELD_ALIASES = {
    "severity": (
        "severity",
        "priority",
        "risk_level",
        "严重程度",
        "风险等级",
        "优先级",
    ),
    "confidence": (
        "confidence",
        "certainty",
        "evidence_confidence",
        "可信度",
        "置信度",
        "证据可信度",
    ),
    "category": (
        "category",
        "issue_type",
        "problem_type",
        "type",
        "问题类型",
        "类型",
        "检查类型",
    ),
    "location": (
        "location",
        "position",
        "where",
        "位置",
        "位置线索",
        "页码",
        "章节",
        "图片位置",
    ),
    "excerpt": (
        "excerpt",
        "quote",
        "original",
        "evidence",
        "document_a_evidence",
        "document_b_evidence",
        "原文摘录",
        "原文",
        "证据",
        "文档A证据",
        "文档B证据",
        "文档线索",
        "图片可见证据",
        "可见线索",
    ),
    "description": (
        "description",
        "issue",
        "problem",
        "finding",
        "疑似问题",
        "问题描述",
        "偏差说明",
        "差异说明",
        "冲突或缺失说明",
        "问题判断",
    ),
    "impact": ("impact", "risk", "影响", "影响说明", "客户影响", "可能影响"),
    "suggestion": (
        "suggestion",
        "recommendation",
        "fix",
        "修改建议",
        "建议修改",
        "建议处理方式",
        "需核对的依据",
    ),
}
REPORT_NO_ACTION_IMPACT_MARKERS = (
    "无实质影响",
    "无实际影响",
    "没有实质影响",
    "没有实际影响",
    "无明显影响",
    "无明确影响",
    "没有明确影响",
    "未发现明确影响",
    "不适用",
    "不造成实质影响",
    "影响不大",
    "影响较小",
    "影响很小",
    "无影响",
    "不影响理解",
    "不影响使用",
    "nosubstantiveimpact",
    "nomaterialimpact",
    "nosignificantimpact",
    "norealimpact",
    "noactualimpact",
    "notapplicable",
    "doesnotaffect",
)
REPORT_NO_ACTION_SUGGESTION_MARKERS = (
    "无需修改",
    "无须修改",
    "不需修改",
    "不需要修改",
    "无需处理",
    "无须处理",
    "不需处理",
    "不需要处理",
    "无需调整",
    "无须调整",
    "不需调整",
    "无需修正",
    "保持不变",
    "nomodificationrequired",
    "nomodificationneeded",
    "noneedtomodify",
    "noneedtochange",
    "nochangeneeded",
    "noactionrequired",
    "noactionneeded",
)
REPORT_LEGACY_LABEL_FIELDS = {
    "问题类型": "category",
    "类型": "category",
    "对象类型": "category",
    "位置": "location",
    "位置线索": "location",
    "图片名称或位置": "location",
    "图片位置": "location",
    "文档线索": "location",
    "原文": "excerpt",
    "原文摘录": "excerpt",
    "文档A证据": "excerpt",
    "文档B证据": "excerpt",
    "资料表述": "excerpt",
    "冲突表述": "excerpt",
    "图片可见内容": "excerpt",
    "图片可见证据": "excerpt",
    "可见内容线索": "excerpt",
    "可见线索": "excerpt",
    "识别到的文字": "excerpt",
    "问题": "description",
    "问题描述": "description",
    "疑似问题": "description",
    "偏差说明": "description",
    "差异说明": "description",
    "冲突或缺失说明": "description",
    "问题判断": "description",
    "不匹配原因": "description",
    "理由": "description",
    "影响": "impact",
    "影响说明": "impact",
    "客户影响": "impact",
    "可能影响": "impact",
    "建议": "suggestion",
    "修改建议": "suggestion",
    "建议修改": "suggestion",
    "建议处理方式": "suggestion",
    "建议补充的标题形式": "suggestion",
}
