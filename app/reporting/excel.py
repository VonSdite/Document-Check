import io
import json

from flask import current_app
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.contracts.task_types import (
    DOCUMENT_TASK_TYPE,
    VIDEO_TASK_TYPE,
    task_type_label,
)
from app.persistence.connection import get_db, now_text
from app.reporting.constants import (
    REPORT_ACCEPTANCE_STATUSES,
    REPORT_EXPORT_EDITABLE_FILL,
    REPORT_EXPORT_HEADER_FILL,
    REPORT_EXPORT_HEADER_FONT,
    REPORT_EXPORT_ITEM_ID_HEADER,
    REPORT_EXPORT_RESULT_CODE_HEADER,
    REPORT_EXPORT_SHEET_NAME,
    REPORT_IMPORT_MAX_ROWS,
    REPORT_ITEM_TYPE_LABEL,
    REPORT_ITEM_TYPES,
    REPORT_REJECTION_REASONS,
    REPORT_TOTAL_EXPORT_ROWS,
    ReportExcelImportError,
)
from app.reporting.service import (
    _apply_report_item_review,
    _cache_prepared_task_report_stats,
    _prepare_task_results,
    _raw_task_results,
    _report_item_fields_for_task,
    _report_item_totals,
    _result_report_items,
    _row_value,
    _task_document_groups,
    _task_results,
)


def build_report_workbook(task):
    results = _task_results(task)
    report_totals = _report_item_totals(results)
    document_groups = _task_document_groups(task)
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = REPORT_EXPORT_SHEET_NAME

    _fill_report_items_sheet(report_sheet, task, results, document_groups)
    _fill_report_totals_sheet(workbook.create_sheet("统计"), report_totals)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def _load_report_excel_reviews(task, payload: bytes) -> int:
    try:
        workbook = load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        current_app.logger.warning(
            "打开报告回填文件失败 task_id=%s error=%s", task["id"], exc
        )
        raise ReportExcelImportError(
            "无法读取回填文件，请上传系统导出的有效 xlsx 报告。"
        ) from exc

    try:
        if REPORT_EXPORT_SHEET_NAME not in workbook.sheetnames:
            raise ReportExcelImportError(
                f"回填文件缺少“{REPORT_EXPORT_SHEET_NAME}”工作表。"
            )
        reviews = _parse_report_excel_reviews(task, workbook[REPORT_EXPORT_SHEET_NAME])
    except ReportExcelImportError:
        raise
    except Exception as exc:
        current_app.logger.warning(
            "解析报告回填文件失败 task_id=%s error=%s", task["id"], exc
        )
        raise ReportExcelImportError(
            "无法读取回填文件，请上传系统导出的有效 xlsx 报告。"
        ) from exc
    finally:
        workbook.close()

    results = _raw_task_results(task)
    result_targets = _report_excel_item_targets(results)
    db = get_db()
    for review in reviews:
        result = result_targets[(review["result_code"], review["item_id"])]
        _apply_report_item_review(
            db,
            task=task,
            result_code=review["result_code"],
            result=result,
            item_id=review["item_id"],
            item_type=review["item_type"],
            acceptance_supplied=review["acceptance_supplied"],
            acceptance_status=review["acceptance_status"],
            rejection_reason=review["rejection_reason"],
            rejection_note=review["rejection_note"],
        )
    task_updated_at = now_text()
    db.execute(
        "UPDATE tasks SET result_json = ?, updated_at = ? WHERE id = ?",
        (json.dumps(results, ensure_ascii=False), task_updated_at, task["id"]),
    )
    db.commit()
    prepared = _prepare_task_results(
        results,
        task_type=task["task_type"] or DOCUMENT_TASK_TYPE,
        task_id=task["id"],
    )
    _cache_prepared_task_report_stats(task, prepared, task_updated_at)
    return len(reviews)


def _parse_report_excel_reviews(task, sheet) -> list[dict]:
    if sheet.max_row > REPORT_IMPORT_MAX_ROWS + 1:
        raise ReportExcelImportError(
            f"回填文件最多允许 {REPORT_IMPORT_MAX_ROWS} 行报告条目。"
        )

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = {}
    for index, value in enumerate(header_row):
        header = _excel_import_text(value)
        if not header:
            continue
        if header in columns:
            raise ReportExcelImportError(
                f"“{REPORT_EXPORT_SHEET_NAME}”工作表存在重复列“{header}”。"
            )
        columns[header] = index

    for required_header in ("任务ID", REPORT_ITEM_TYPE_LABEL):
        if required_header not in columns:
            raise ReportExcelImportError(f"回填文件缺少“{required_header}”列。")

    metadata_headers = (REPORT_EXPORT_RESULT_CODE_HEADER, REPORT_EXPORT_ITEM_ID_HEADER)
    has_metadata = all(header in columns for header in metadata_headers)
    if any(header in columns for header in metadata_headers) and not has_metadata:
        raise ReportExcelImportError("回填文件的隐藏条目标识列不完整，请重新导出报告。")
    if not has_metadata and not all(header in columns for header in ("检查项", "条目")):
        raise ReportExcelImportError("回填文件缺少条目标识列，无法匹配报告条目。")

    results = _raw_task_results(task)
    item_targets = _report_excel_item_targets(results)
    legacy_targets = _legacy_report_excel_item_targets(task, results)
    type_values = {
        **{code: code for code in REPORT_ITEM_TYPES},
        **{label: code for code, label in REPORT_ITEM_TYPES.items()},
    }
    acceptance_values = {
        **{code: code for code in REPORT_ACCEPTANCE_STATUSES},
        **{label: code for code, label in REPORT_ACCEPTANCE_STATUSES.items()},
    }
    rejection_values = {
        **{code: code for code in REPORT_REJECTION_REASONS},
        **{label: code for code, label in REPORT_REJECTION_REASONS.items()},
    }

    reviews = []
    seen_targets = set()
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        item_marker = _excel_import_text(_excel_row_value(row, columns, "条目"))
        metadata_item_id = _excel_import_text(
            _excel_row_value(row, columns, REPORT_EXPORT_ITEM_ID_HEADER)
        )
        if not item_marker and not metadata_item_id:
            continue

        item_type_text = _excel_import_text(
            _excel_row_value(row, columns, REPORT_ITEM_TYPE_LABEL)
        )
        if not item_type_text:
            continue
        item_type = type_values.get(item_type_text)
        if item_type is None:
            raise ReportExcelImportError(
                f"第 {row_number} 行“{REPORT_ITEM_TYPE_LABEL}”无效，请选择问题、建议或非问题。"
            )

        task_id = _excel_import_task_id(_excel_row_value(row, columns, "任务ID"))
        if task_id != str(task["id"]):
            raise ReportExcelImportError(f"第 {row_number} 行任务ID与当前报告不一致。")

        if has_metadata:
            result_code = _excel_import_text(
                _excel_row_value(row, columns, REPORT_EXPORT_RESULT_CODE_HEADER)
            )
            item_id = metadata_item_id
            if not result_code or not item_id:
                raise ReportExcelImportError(
                    f"第 {row_number} 行条目标识缺失，请重新导出报告。"
                )
            target_key = (result_code, item_id)
        else:
            result_name = _excel_import_text(_excel_row_value(row, columns, "检查项"))
            candidates = legacy_targets.get((result_name, item_marker), [])
            if len(candidates) != 1:
                raise ReportExcelImportError(
                    f"第 {row_number} 行无法唯一匹配当前报告条目，请重新导出报告后标注。"
                )
            target_key = candidates[0]

        if target_key not in item_targets:
            raise ReportExcelImportError(
                f"第 {row_number} 行报告条目不存在或已发生变化，请重新导出报告。"
            )
        if target_key in seen_targets:
            raise ReportExcelImportError(
                f"第 {row_number} 行与前面的行重复指向同一报告条目。"
            )
        seen_targets.add(target_key)

        acceptance_supplied = False
        acceptance_status = None
        rejection_reason = ""
        rejection_note = ""
        if "是否接纳" in columns:
            acceptance_text = _excel_import_text(
                _excel_row_value(row, columns, "是否接纳")
            )
            if acceptance_text:
                acceptance_status = acceptance_values.get(acceptance_text)
                if acceptance_status is None:
                    raise ReportExcelImportError(
                        f"第 {row_number} 行“是否接纳”无效，请选择未确认、接纳或不接纳。"
                    )
                acceptance_supplied = True
                if acceptance_status == "rejected":
                    reason_text = _excel_import_text(
                        _excel_row_value(row, columns, "不接纳原因")
                    )
                    if not reason_text:
                        raise ReportExcelImportError(
                            f"第 {row_number} 行选择不接纳时必须填写“不接纳原因”。"
                        )
                    rejection_reason = rejection_values.get(reason_text, "")
                    if not rejection_reason:
                        raise ReportExcelImportError(
                            f"第 {row_number} 行“不接纳原因”无效。"
                        )
                    rejection_note = _excel_import_text(
                        _excel_row_value(row, columns, "人工原因")
                    )
                    if rejection_reason == "other" and not rejection_note:
                        raise ReportExcelImportError(
                            f"第 {row_number} 行选择其他原因时必须填写人工原因。"
                        )

        reviews.append(
            {
                "result_code": target_key[0],
                "item_id": target_key[1],
                "item_type": item_type,
                "acceptance_supplied": acceptance_supplied,
                "acceptance_status": acceptance_status,
                "rejection_reason": rejection_reason,
                "rejection_note": rejection_note,
            }
        )

    if not reviews:
        raise ReportExcelImportError("回填文件中没有可识别的报告标注。")
    return reviews


def _report_excel_item_targets(results: list[dict]) -> dict[tuple[str, str], dict]:
    targets = {}
    for result in results:
        result_code = str(result.get("code") or "")
        for item in _result_report_items(result):
            targets[(result_code, str(item.get("id") or ""))] = result
    return targets


def _legacy_report_excel_item_targets(
    task, results: list[dict]
) -> dict[tuple[str, str], list[tuple[str, str]]]:
    prepared = _prepare_task_results(
        results,
        task_type=task["task_type"] or DOCUMENT_TASK_TYPE,
        task_id=task["id"],
    )
    targets = {}
    for result in prepared:
        result_name = str(result.get("name") or "").strip()
        result_code = str(result.get("code") or "")
        for item in result.get("report_items") or []:
            key = (result_name, f"条目 {item.get('index')}")
            targets.setdefault(key, []).append((result_code, str(item.get("id") or "")))
    return targets


def _excel_row_value(row: tuple, columns: dict[str, int], header: str):
    index = columns.get(header)
    if index is None or index >= len(row):
        return None
    return row[index]


def _excel_import_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _excel_import_task_id(value) -> str:
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _excel_import_text(value)


def _fill_report_items_sheet(
    sheet, task, results: list[dict], document_groups: list[dict]
) -> None:
    report_item_fields = _report_item_fields_for_task(task["task_type"])
    headers = [
        "任务ID",
        "任务类型",
        "文件名称",
        "检查项",
        "条目",
        *[label for _, label in report_item_fields],
        REPORT_ITEM_TYPE_LABEL,
        "是否接纳",
        "不接纳原因",
        "人工原因",
        REPORT_EXPORT_RESULT_CODE_HEADER,
        REPORT_EXPORT_ITEM_ID_HEADER,
    ]
    sheet.append(headers)
    context = _excel_task_context(task, document_groups)
    for result in results:
        report_items = result.get("report_items") or []
        if not report_items:
            sheet.append(
                [
                    *context,
                    _excel_cell_text(result.get("name")),
                    "",
                    *["" for _ in report_item_fields],
                    "未拆分",
                    "",
                    "",
                    "",
                    _excel_cell_text(result.get("code")),
                    "",
                ]
            )
            continue
        for item in report_items:
            sheet.append(
                [
                    *context,
                    _excel_cell_text(result.get("name")),
                    f"条目 {item.get('index')}",
                    *[
                        _excel_report_item_field(item, field, task["task_type"])
                        for field, _ in report_item_fields
                    ],
                    _excel_cell_text(item.get("type_label")),
                    _excel_cell_text(item.get("acceptance_label")),
                    _excel_cell_text(item.get("rejection_reason_label"))
                    if item.get("acceptance_status") == "rejected"
                    else "",
                    _excel_cell_text(item.get("rejection_note"))
                    if item.get("acceptance_status") == "rejected"
                    else "",
                    _excel_cell_text(result.get("code")),
                    _excel_cell_text(item.get("id")),
                ]
            )
    _style_excel_sheet(sheet)
    _configure_report_review_sheet(sheet, headers)


def _excel_report_item_field(item: dict, field: str, task_type: str | None):
    value = _excel_cell_text(item.get(field))
    if (task_type or DOCUMENT_TASK_TYPE) != VIDEO_TASK_TYPE or field != "excerpt":
        return value
    evidence_lines = []
    for ref in item.get("evidence_refs") or []:
        position = str(ref.get("position") or "").strip()
        filename = str(ref.get("filename") or "").strip()
        label = position
        if filename:
            label = f"{label}（{filename}）" if label else filename
        if label and label not in evidence_lines:
            evidence_lines.append(label)
    if not evidence_lines:
        return value
    evidence_text = "关键帧：" + "、".join(evidence_lines)
    return f"{value}\n{evidence_text}" if value else evidence_text


def _configure_report_review_sheet(sheet, headers: list[str]) -> None:
    columns = {header: index + 1 for index, header in enumerate(headers)}
    last_row = max(sheet.max_row, 2)
    validation_choices = (
        (REPORT_ITEM_TYPE_LABEL, tuple(REPORT_ITEM_TYPES.values())),
        ("是否接纳", tuple(REPORT_ACCEPTANCE_STATUSES.values())),
        ("不接纳原因", tuple(REPORT_REJECTION_REASONS.values())),
    )
    for header, choices in validation_choices:
        column = columns[header]
        column_letter = get_column_letter(column)
        validation = DataValidation(
            type="list",
            formula1=f'"{",".join(choices)}"',
            allow_blank=header == "不接纳原因",
        )
        validation.error = f"请从下拉列表中选择有效的{header}。"
        validation.errorTitle = "标注值无效"
        validation.prompt = f"请选择{header}。"
        validation.promptTitle = "报告标注"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{column_letter}2:{column_letter}{last_row}")
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=column).fill = REPORT_EXPORT_EDITABLE_FILL

    for header in (REPORT_EXPORT_RESULT_CODE_HEADER, REPORT_EXPORT_ITEM_ID_HEADER):
        sheet.column_dimensions[get_column_letter(columns[header])].hidden = True
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"


def _fill_report_totals_sheet(sheet, report_totals: dict) -> None:
    sheet.append(["指标", "值"])
    for label, key in REPORT_TOTAL_EXPORT_ROWS:
        sheet.append([label, report_totals.get(key, 0)])
    _style_excel_sheet(sheet)


def _excel_task_context(task, document_groups: list[dict]) -> list:
    return [
        _row_value(task, "id", ""),
        task_type_label(_row_value(task, "task_type", DOCUMENT_TASK_TYPE)),
        _excel_document_names(task, document_groups),
    ]


def _excel_document_names(task, document_groups: list[dict]) -> str:
    names = []
    for group in document_groups:
        for file in group.get("files", []):
            name = str(file.get("original_filename") or "").strip()
            if name:
                names.append(name)
    if names:
        return "\n".join(names)
    return _excel_cell_text(_row_value(task, "original_filename", ""))


def _excel_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float)):
        return value
    return str(value).strip()


def _style_excel_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = REPORT_EXPORT_HEADER_FONT
        cell.fill = REPORT_EXPORT_HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells[:200]) + 2
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(width, 10), 60
        )
