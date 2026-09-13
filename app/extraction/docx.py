import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .common import (
    _clean_hyperlink_target,
    _format_hyperlink_text,
    _record_hyperlink,
)

_HYPERLINK_FIELD_PATTERN = re.compile(
    r"\bHYPERLINK\s+(?:(\\l)\s+)?(?:\"([^\"]+)\"|(\S+))",
    re.IGNORECASE,
)


def _extract_docx(path: Path, hyperlinks: list[dict] | None = None) -> str:
    document = Document(str(path))
    internal_targets = {
        str(node.get(qn("w:name")) or "").strip()
        for node in document.element.iter(qn("w:bookmarkStart"))
        if str(node.get(qn("w:name")) or "").strip()
    }
    parts = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = _docx_paragraph_text(
            paragraph,
            hyperlinks,
            location=f"第{paragraph_index}段",
            internal_targets=internal_targets,
        )
        if text:
            parts.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for column_index, cell in enumerate(row.cells, start=1):
                location = f"表格{table_index} > 第{row_index}行 > 第{column_index}列"
                cell_text = "\n".join(
                    text
                    for paragraph in cell.paragraphs
                    if (
                        text := _docx_paragraph_text(
                            paragraph,
                            hyperlinks,
                            location=location,
                            internal_targets=internal_targets,
                        )
                    )
                ).strip()
                if cell_text:
                    cells.append(cell_text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _docx_paragraph_text(
    paragraph,
    hyperlinks: list[dict] | None = None,
    *,
    location: str = "",
    internal_targets: set[str] | None = None,
) -> str:
    text = _docx_element_text(
        paragraph._p,
        paragraph.part,
        hyperlinks,
        location,
        internal_targets or set(),
    ).strip()
    for target in _docx_complex_field_hyperlinks(paragraph._p):
        _record_hyperlink(
            hyperlinks,
            "",
            target,
            location,
            internal_target_exists=(
                target.lstrip("#") in (internal_targets or set())
                if target.startswith("#")
                else None
            ),
        )
        if target not in text:
            text += f"（超链接：{target}）"
    return text


def _docx_element_text(
    element,
    part,
    hyperlinks: list[dict] | None,
    location: str,
    internal_targets: set[str],
) -> str:
    if element.tag == qn("w:hyperlink"):
        label = _docx_plain_element_text(element)
        target = _docx_hyperlink_target(element, part)
        _record_hyperlink(
            hyperlinks,
            label,
            target,
            location,
            internal_target_exists=(
                target.lstrip("#") in internal_targets
                if target.startswith("#")
                else None
            ),
        )
        return _format_hyperlink_text(label, target)
    if element.tag == qn("w:fldSimple"):
        label = _docx_plain_element_text(element)
        target = _docx_field_hyperlink_target(element.get(qn("w:instr")))
        _record_hyperlink(
            hyperlinks,
            label,
            target,
            location,
            internal_target_exists=(
                target.lstrip("#") in internal_targets
                if target.startswith("#")
                else None
            ),
        )
        return _format_hyperlink_text(label, target)
    if element.tag == qn("w:t"):
        return element.text or ""
    if element.tag == qn("w:tab"):
        return "\t"
    if element.tag in {qn("w:br"), qn("w:cr")}:
        return "\n"
    if element.tag == qn("w:noBreakHyphen"):
        return "-"
    if element.tag in {qn("w:instrText"), qn("w:delText")}:
        return ""
    return "".join(
        _docx_element_text(child, part, hyperlinks, location, internal_targets)
        for child in element
    )


def _docx_plain_element_text(element) -> str:
    if element.tag == qn("w:t"):
        return element.text or ""
    if element.tag == qn("w:tab"):
        return "\t"
    if element.tag in {qn("w:br"), qn("w:cr")}:
        return "\n"
    if element.tag == qn("w:noBreakHyphen"):
        return "-"
    if element.tag in {qn("w:instrText"), qn("w:delText")}:
        return ""
    return "".join(_docx_plain_element_text(child) for child in element)


def _docx_hyperlink_target(element, part) -> str:
    relationship_id = element.get(qn("r:id"))
    if relationship_id:
        try:
            return _clean_hyperlink_target(part.rels[relationship_id].target_ref)
        except (AttributeError, KeyError):
            pass
    anchor = _clean_hyperlink_target(element.get(qn("w:anchor")))
    return f"#{anchor}" if anchor else ""


def _docx_complex_field_hyperlinks(element) -> list[str]:
    instruction = "".join(node.text or "" for node in element.iter(qn("w:instrText")))
    return _hyperlink_targets_from_field_instruction(instruction)


def _docx_field_hyperlink_target(instruction) -> str:
    targets = _hyperlink_targets_from_field_instruction(instruction)
    return targets[0] if targets else ""


def _hyperlink_targets_from_field_instruction(instruction) -> list[str]:
    targets = []
    for match in _HYPERLINK_FIELD_PATTERN.finditer(str(instruction or "")):
        target = _clean_hyperlink_target(match.group(2) or match.group(3))
        if match.group(1) and target:
            target = f"#{target.lstrip('#')}"
        if target and target not in targets:
            targets.append(target)
    return targets
