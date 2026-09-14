import heapq
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet._reader import FORMULA_TAG, WorkSheetParser

from app.documents.extraction.common import (
    _clean_hyperlink_target,
    _format_hyperlink_text,
    _record_hyperlink,
    check_extraction_canceled,
)

_SPREADSHEETML_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

_OFFICE_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)


def _extract_openpyxl_workbook(
    path: Path,
    extracted_hyperlinks: list[dict] | None = None,
    *,
    cancel_event=None,
) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        check_extraction_canceled(cancel_event)
        try:
            hyperlinks = _extract_xlsx_hyperlinks(
                path, workbook, cancel_event=cancel_event
            )
        except (ET.ParseError, KeyError, OSError, zipfile.BadZipFile):
            hyperlinks = {}
        parts = []
        workbook_titles = set(workbook.sheetnames)
        for sheet in workbook.worksheets:
            check_extraction_canceled(cancel_event)
            rows = _openpyxl_sheet_rows_text(
                sheet,
                hyperlinks.get(sheet.title, []),
                extracted_hyperlinks,
                workbook_titles,
                cancel_event=cancel_event,
            )
            if rows:
                parts.append(f"# 工作表：{sheet.title}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    finally:
        workbook.close()


def _extract_xls(
    path: Path, hyperlinks: list[dict] | None = None, *, cancel_event=None
) -> str:
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        parts = []
        for sheet in workbook.sheets():
            rows = []
            for row_index in range(sheet.nrows):
                check_extraction_canceled(cancel_event)
                values = []
                for column_index in range(sheet.ncols):
                    value = sheet.cell_value(row_index, column_index)
                    hyperlink = sheet.hyperlink_map.get((row_index, column_index))
                    if hyperlink is not None:
                        target = _clean_hyperlink_target(hyperlink.url_or_path)
                        textmark = _clean_hyperlink_target(hyperlink.textmark)
                        if textmark:
                            target = (
                                f"{target}#{textmark}" if target else f"#{textmark}"
                            )
                        label = value or hyperlink.desc
                        value = _format_hyperlink_text(label, target)
                        _record_hyperlink(
                            hyperlinks,
                            label,
                            target,
                            (
                                f"工作表“{sheet.name}”!"
                                f"{get_column_letter(column_index + 1)}{row_index + 1}"
                            ),
                            internal_target_exists=(
                                _spreadsheet_internal_target_exists(
                                    target,
                                    set(workbook.sheet_names()),
                                )
                                if target.startswith("#")
                                else None
                            ),
                        )
                    values.append(value)
                row_text = _spreadsheet_row_text(values)
                if row_text:
                    rows.append(row_text)
            if rows:
                parts.append(f"# 工作表：{sheet.name}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    finally:
        workbook.release_resources()


class _WorksheetValuesAndFormulas(WorkSheetParser):
    """在同一单元格解析中保留缓存值与公式，使用 openpyxl 的类型和共享公式规则。"""

    def parse_cell(self, element):
        cell = super().parse_cell(element)
        if element.find(FORMULA_TAG) is not None:
            cell["formula"] = self.parse_formula(element)
        return cell


class _WorksheetHyperlinks:
    """按升序行定位活动链接范围，保留 XML 中的声明顺序。"""

    def __init__(self, hyperlinks):
        self.starts = sorted(
            (item[1], index, item) for index, item in enumerate(hyperlinks)
        )
        self.position = 0
        self.active = {}
        self.ends = []

    def for_row(self, row_index):
        while (
            self.position < len(self.starts)
            and self.starts[self.position][0] <= row_index
        ):
            _, index, item = self.starts[self.position]
            self.position += 1
            if item[3] >= row_index:
                self.active[index] = item
                heapq.heappush(self.ends, (item[3], index))
        while self.ends and self.ends[0][0] < row_index:
            _, index = heapq.heappop(self.ends)
            self.active.pop(index, None)
        return _RowHyperlinks(self.active.items())

    def rows_through(self, first_row, last_row):
        """遍历含链接的缺省行与当前数据行，跳过连续空白区域。"""
        row_index = first_row
        while row_index <= last_row:
            row_links = self.for_row(row_index)
            if row_links.last_column or row_index == last_row:
                yield row_index, row_links
            if row_index == last_row:
                break
            if self.active:
                row_index += 1
            elif self.position < len(self.starts):
                row_index = min(last_row, self.starts[self.position][0])
            else:
                row_index = last_row


class _RowHyperlinks:
    """按升序列访问链接，重叠范围使用最先声明的目标。"""

    def __init__(self, hyperlinks):
        self.starts = sorted(
            (item[0], index, item[2], item[4], item[5]) for index, item in hyperlinks
        )
        self.last_column = max((item[2] for item in self.starts), default=0)
        self.position = 0
        self.active = []

    def at(self, column_index):
        while (
            self.position < len(self.starts)
            and self.starts[self.position][0] <= column_index
        ):
            _, index, end, target, display = self.starts[self.position]
            heapq.heappush(self.active, (index, end, target, display))
            self.position += 1
        while self.active and self.active[0][1] < column_index:
            heapq.heappop(self.active)
        return self.active[0][2:] if self.active else None


def _openpyxl_sheet_rows_text(
    sheet,
    hyperlinks,
    extracted_hyperlinks: list[dict] | None = None,
    workbook_titles: set[str] | None = None,
    *,
    cancel_event=None,
) -> list[str]:
    rows = []
    links = _WorksheetHyperlinks(hyperlinks)
    max_row, max_column = sheet.max_row, sheet.max_column
    with sheet._get_source() as source:
        parser = _WorksheetValuesAndFormulas(
            source,
            sheet._shared_strings,
            data_only=True,
            epoch=sheet.parent.epoch,
            date_formats=sheet.parent._date_formats,
            timedelta_formats=sheet.parent._timedelta_formats,
        )
        previous_row = 0
        for row_index, cells in parser.parse():
            if row_index <= previous_row:
                continue
            if max_row and row_index > max_row:
                break
            # 缺省行只处理其链接，单元格按实际内容与链接范围展开。
            for index, row_links in links.rows_through(previous_row + 1, row_index):
                check_extraction_canceled(cancel_event)
                values = {}
                formulas = {}
                if index == row_index:
                    for cell in cells:
                        column = cell["column"]
                        if max_column and column > max_column:
                            continue
                        value = cell["value"]
                        formula = cell.get("formula")
                        if (
                            value is None
                            and isinstance(formula, str)
                            and formula.startswith("=")
                        ):
                            value = formula
                        if value is not None:
                            values[column] = value
                        else:
                            values.pop(column, None)
                        if formula is not None:
                            formulas[column] = formula
                        else:
                            formulas.pop(column, None)
                last_column = max(max(values, default=0), row_links.last_column)
                if max_column:
                    last_column = min(max_column, last_column)
                if not last_column:
                    continue
                output = []
                for column in range(1, last_column + 1):
                    value = values.get(column)
                    hyperlink = row_links.at(column)
                    target = _xlsx_formula_hyperlink_target(
                        formulas.get(column, value)
                    ) or (hyperlink[0] if hyperlink else "")
                    if target:
                        label = value or (hyperlink[1] if hyperlink else "")
                        value = _format_hyperlink_text(label, target)
                        _record_hyperlink(
                            extracted_hyperlinks,
                            label,
                            target,
                            f"工作表“{sheet.title}”!{get_column_letter(column)}{index}",
                            internal_target_exists=(
                                _spreadsheet_internal_target_exists(
                                    target, workbook_titles or set()
                                )
                                if target.startswith("#")
                                else None
                            ),
                        )
                    output.append(value)
                row_text = _spreadsheet_row_text(output)
                if row_text:
                    rows.append(row_text)
            previous_row = row_index
    return rows


def _extract_xlsx_hyperlinks(
    path: Path, workbook, *, cancel_event=None
) -> dict[str, list[tuple]]:
    result = {}
    with zipfile.ZipFile(path) as archive:
        archive_names = set(archive.namelist())
        for sheet in workbook.worksheets:
            sheet_path = str(getattr(sheet, "_worksheet_path", "")).replace("\\", "/")
            if not sheet_path or sheet_path not in archive_names:
                continue
            relationships = _xlsx_sheet_relationships(
                archive,
                archive_names,
                sheet_path,
            )
            items = []
            with archive.open(sheet_path) as source:
                for _event, element in ET.iterparse(source, events=("end",)):
                    if element.tag == f"{{{_SPREADSHEETML_NAMESPACE}}}row":
                        check_extraction_canceled(cancel_event)
                    if element.tag != f"{{{_SPREADSHEETML_NAMESPACE}}}hyperlink":
                        element.clear()
                        continue
                    cell_range = str(element.get("ref") or "").strip()
                    relationship_id = element.get(
                        f"{{{_OFFICE_RELATIONSHIP_NAMESPACE}}}id"
                    )
                    target = _clean_hyperlink_target(
                        relationships.get(relationship_id, "")
                    )
                    if not target:
                        location = _clean_hyperlink_target(element.get("location"))
                        target = f"#{location}" if location else ""
                    try:
                        bounds = range_boundaries(cell_range)
                    except ValueError:
                        bounds = None
                    if target and bounds:
                        items.append((*bounds, target, element.get("display") or ""))
                    element.clear()
            if items:
                result[sheet.title] = items
    return result


def _xlsx_sheet_relationships(
    archive: zipfile.ZipFile,
    archive_names: set[str],
    sheet_path: str,
) -> dict[str, str]:
    sheet_name = sheet_path.rsplit("/", 1)[-1]
    sheet_dir = sheet_path.rsplit("/", 1)[0]
    relationship_path = f"{sheet_dir}/_rels/{sheet_name}.rels"
    if relationship_path not in archive_names:
        return {}
    relationships = {}
    with archive.open(relationship_path) as source:
        for _event, element in ET.iterparse(source, events=("end",)):
            if element.tag.rsplit("}", 1)[-1] == "Relationship":
                relationship_id = element.get("Id")
                target = _clean_hyperlink_target(element.get("Target"))
                if relationship_id and target:
                    relationships[relationship_id] = target
            element.clear()
    return relationships


def _xlsx_formula_hyperlink_target(formula) -> str:
    if not isinstance(formula, str) or not formula.startswith("="):
        return ""
    match = re.match(
        r'^=HYPERLINK\(\s*"((?:[^"]|"")*)"',
        formula,
        re.IGNORECASE,
    )
    if not match:
        return ""
    return _clean_hyperlink_target(match.group(1).replace('""', '"'))


def _spreadsheet_internal_target_exists(target: str, workbook_titles: set[str]) -> bool:
    location = str(target or "").lstrip("#").strip()
    if not location:
        return False
    sheet_name = location.split("!", 1)[0].strip().strip("'").replace("''", "'")
    if "!" not in location:
        return True
    return sheet_name in workbook_titles


def _spreadsheet_rows_text(rows) -> list[str]:
    result = []
    for row in rows:
        row_text = _spreadsheet_row_text(row)
        if row_text:
            result.append(row_text)
    return result


def _spreadsheet_row_text(values) -> str:
    cells = [_spreadsheet_cell_text(value) for value in values]
    while cells and not cells[-1]:
        cells.pop()
    return " | ".join(cells)


def _spreadsheet_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
