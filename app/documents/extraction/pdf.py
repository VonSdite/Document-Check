import html
import logging
import re
import unicodedata
from bisect import bisect_left, bisect_right
from pathlib import Path

import fitz
from openpyxl.utils.cell import get_column_letter
from pypdf import PdfReader

from app.documents.extraction.common import (
    DocumentReadCanceled,
    _clean_hyperlink_target,
    _format_hyperlink_text,
    _record_hyperlink,
    check_extraction_canceled,
)

_PDF_TABLE_COORDINATE_TOLERANCE = 2.0
_PDF_TABLE_MIN_TEXT_COVERAGE = 0.85
logger = logging.getLogger(__name__)


logger.addHandler(logging.NullHandler())


def _extract_pdf(
    path: Path,
    extracted_hyperlinks: list[dict] | None = None,
    *,
    include_tables: bool = True,
    cancel_event=None,
) -> str:
    reader = None
    try:
        layout_document = fitz.open(str(path))
    except Exception:
        layout_document = None
    if layout_document is None:
        reader = PdfReader(str(path))
    try:
        pages = []
        page_count = (
            layout_document.page_count
            if layout_document is not None
            else len(reader.pages)
        )
        for index in range(1, page_count + 1):
            check_extraction_canceled(cancel_event)
            pypdf_page = None
            if layout_document is not None and index <= layout_document.page_count:
                layout_page = layout_document[index - 1]
                pymupdf_text, raw_page = _extract_pymupdf_page_text(layout_page)
                pypdf_text = ""
                if _pdf_should_try_pypdf(pymupdf_text):
                    if reader is None:
                        reader = PdfReader(str(path))
                    pypdf_page = reader.pages[index - 1]
                    pypdf_text = pypdf_page.extract_text() or ""
                pypdf_text = _normalize_pdf_overlapping_spaces(pypdf_text, raw_page)
                pymupdf_text = _normalize_pdf_overlapping_spaces(pymupdf_text, raw_page)
                text = _select_pdf_page_text(pypdf_text, pymupdf_text)
                if include_tables and text == pymupdf_text:
                    structured_text = _extract_pymupdf_page_with_tables(
                        layout_page,
                        index,
                        raw_page,
                        pymupdf_text,
                        cancel_event=cancel_event,
                    )
                    if structured_text:
                        text = structured_text
                hyperlinks = _extract_pymupdf_page_hyperlinks(layout_page)
            else:
                pypdf_page = reader.pages[index - 1]
                pypdf_text = pypdf_page.extract_text() or ""
                text = pypdf_text
                hyperlinks = _extract_pypdf_page_hyperlinks(pypdf_page)
            if hyperlinks:
                text = _append_hyperlinks(text, hyperlinks)
                for label, target, internal_target_exists in hyperlinks:
                    _record_hyperlink(
                        extracted_hyperlinks,
                        label,
                        target,
                        f"第{index}页",
                        internal_target_exists=internal_target_exists,
                    )
            if text.strip():
                pages.append(f"[第{index}页]\n{text.strip()}")
        return "\n\n".join(pages)
    finally:
        if layout_document is not None:
            layout_document.close()


def _pdf_should_try_pypdf(text: str) -> bool:
    if not str(text or "").strip():
        return True
    return _pdf_text_quality(text)["bad_char_ratio"] >= 0.01


def _extract_pymupdf_page_text(page) -> tuple[str, dict]:
    try:
        text_page = page.get_textpage()
        text = page.get_text("text", textpage=text_page, sort=False) or ""
        raw_page = page.get_text("rawdict", textpage=text_page)
        return text, raw_page
    except Exception:
        return "", {}


def _extract_pymupdf_page_with_tables(
    page,
    page_number: int,
    raw_page: dict,
    source_text: str,
    *,
    cancel_event=None,
) -> str:
    try:
        drawings = page.get_drawings()
        if not _pdf_has_table_edges(drawings):
            return ""
        image_bboxes = [
            tuple(float(value) for value in image.get("bbox"))
            for image in page.get_image_info()
            if image.get("bbox")
        ]
        table_finder = page.find_tables(strategy="lines", paths=drawings)
        candidate_tables = list(table_finder.tables)
    except Exception as exc:
        logger.debug("PDF 表格检测失败 page=%s error=%s", page_number, exc)
        return ""

    if not candidate_tables:
        return ""
    table_models = []
    text_index = _PdfTextPresenceIndex(raw_page)
    # 单页所有表格的空单元格共用一个文本对象。
    textpage = None
    try:
        textpage = page.get_textpage()
    except Exception:
        pass
    for table_index, table in enumerate(candidate_tables, start=1):
        check_extraction_canceled(cancel_event)
        try:
            table_bbox = tuple(float(value) for value in table.bbox)
        except (AttributeError, TypeError, ValueError):
            table_bbox = None
        if table_bbox is None:
            table_drawings = drawings
            table_images = image_bboxes
        else:
            table_rect = fitz.Rect(table_bbox)
            table_drawings = [
                drawing
                for drawing in drawings
                if _pdf_drawing_intersects(drawing, table_rect)
            ]
            table_images = [
                image_bbox
                for image_bbox in image_bboxes
                if _pdf_rects_intersect(image_bbox, table_bbox)
            ]
        try:
            model = _pdf_table_model(
                page,
                table,
                page_number=page_number,
                table_index=table_index,
                drawings=table_drawings,
                image_bboxes=table_images,
                cancel_event=cancel_event,
                textpage=textpage,
                text_index=text_index,
            )
        except DocumentReadCanceled:
            raise
        except Exception as exc:
            logger.debug(
                "PDF 单个表格结构化失败 page=%s table=%s error=%s",
                page_number,
                table_index,
                exc,
            )
            continue
        if model is not None:
            table_models.append(model)
    if not table_models:
        return ""

    table_rects = [model["bbox"] for model in table_models]
    elements = _pdf_page_text_elements(raw_page, table_rects)
    for model in table_models:
        elements.append(
            {
                "bbox": model["bbox"],
                "content": _format_pdf_table_html(model),
                "plain_text": "\n".join(
                    cell["text"] for cell in model["cells"] if cell["text"]
                ),
                "kind": "table",
                "source_order": model["table_index"],
            }
        )
    elements.sort(key=_pdf_page_element_sort_key)
    plain_text = "\n".join(
        element["plain_text"] for element in elements if element["plain_text"]
    )
    source_chars = _pdf_text_quality(source_text)["useful_chars"]
    structured_chars = _pdf_text_quality(plain_text)["useful_chars"]
    if source_chars and structured_chars < source_chars * _PDF_TABLE_MIN_TEXT_COVERAGE:
        logger.debug(
            "PDF 表格结构化结果文本覆盖不足 page=%s source_chars=%s structured_chars=%s",
            page_number,
            source_chars,
            structured_chars,
        )
        return ""
    return "\n".join(
        element["content"] for element in elements if element["content"]
    ).strip()


def _pdf_has_table_edges(drawings: list[dict]) -> bool:
    edge_count = 0
    for drawing in drawings:
        for item in drawing.get("items", []):
            if item and item[0] in {"l", "re"}:
                edge_count += 1
                if edge_count >= 4:
                    return True
    return False


def _pdf_drawing_intersects(drawing: dict, table_rect: fitz.Rect) -> bool:
    drawing_rect = drawing.get("rect")
    return drawing_rect is not None and fitz.Rect(drawing_rect).intersects(table_rect)


def _pdf_rects_intersect(first: tuple[float, ...], second: tuple[float, ...]) -> bool:
    return not (
        first[2] <= second[0]
        or first[0] >= second[2]
        or first[3] <= second[1]
        or first[1] >= second[3]
    )


def _pdf_table_model(
    page,
    table,
    *,
    page_number: int,
    table_index: int,
    drawings: list[dict],
    image_bboxes: list[tuple[float, ...]],
    cancel_event=None,
    textpage=None,
    text_index=None,
) -> dict | None:
    try:
        row_count = int(table.row_count)
        column_count = int(table.col_count)
        rows = list(table.rows)
        extracted_rows = table.extract()
        table_bbox = tuple(float(value) for value in table.bbox)
    except Exception:
        return None
    if row_count < 2 or column_count < 2 or len(rows) != row_count:
        return None

    row_boundaries = _cluster_pdf_coordinates(
        [
            coordinate
            for row in rows
            for coordinate in (float(row.bbox[1]), float(row.bbox[3]))
        ]
    )
    column_boundaries = _cluster_pdf_coordinates(
        [
            coordinate
            for row in rows
            for cell_bbox in row.cells
            if cell_bbox is not None
            for coordinate in (float(cell_bbox[0]), float(cell_bbox[2]))
        ]
    )
    if (
        len(row_boundaries) != row_count + 1
        or len(column_boundaries) != column_count + 1
    ):
        return None

    nontext = _PdfNontextIndex(image_bboxes, drawings)
    cells = []
    covered_positions = set()
    for row_index, row in enumerate(rows):
        check_extraction_canceled(cancel_event)
        if len(row.cells) != column_count:
            return None
        for column_index, cell_bbox in enumerate(row.cells):
            if cell_bbox is None:
                continue
            bbox = tuple(float(value) for value in cell_bbox)
            start_row = _pdf_coordinate_index(row_boundaries, bbox[1])
            end_row = _pdf_coordinate_index(row_boundaries, bbox[3])
            start_column = _pdf_coordinate_index(column_boundaries, bbox[0])
            end_column = _pdf_coordinate_index(column_boundaries, bbox[2])
            if (
                start_row != row_index
                or start_column != column_index
                or end_row is None
                or end_column is None
                or end_row <= row_index
                or end_column <= column_index
            ):
                return None
            rowspan = end_row - row_index
            colspan = end_column - column_index
            for covered_row in range(row_index, end_row):
                for covered_column in range(column_index, end_column):
                    covered_positions.add((covered_row, covered_column))
            cell_text = ""
            extracted_row = (
                extracted_rows[row_index]
                if isinstance(extracted_rows, (list, tuple))
                and row_index < len(extracted_rows)
                else ()
            )
            if isinstance(extracted_row, (list, tuple)) and column_index < len(
                extracted_row
            ):
                cell_text = str(extracted_row[column_index] or "").strip()
            if not cell_text and (text_index is None or text_index.may_have_text(bbox)):
                try:
                    if textpage is None:
                        textpage = page.get_textpage()
                    cell_text = str(
                        page.get_textbox(fitz.Rect(bbox), textpage=textpage) or ""
                    ).strip()
                except Exception:
                    cell_text = ""
            has_nontext_content = nontext.contains(bbox)
            cells.append(
                {
                    "row": row_index,
                    "column": column_index,
                    "rowspan": rowspan,
                    "colspan": colspan,
                    "bbox": bbox,
                    "text": cell_text,
                    "has_nontext_content": has_nontext_content,
                }
            )

    unresolved_positions = [
        (row_index, column_index)
        for row_index, row in enumerate(rows)
        for column_index, cell_bbox in enumerate(row.cells)
        if cell_bbox is None and (row_index, column_index) not in covered_positions
    ]
    if not cells or not any(cell["text"] for cell in cells):
        return None
    vector_edge_count = _pdf_table_vector_edge_count(drawings, table_bbox)
    if unresolved_positions:
        confidence = "low"
    elif vector_edge_count >= max(4, row_count + column_count):
        confidence = "high"
    else:
        confidence = "medium"
    return {
        "id": f"page{page_number:03d}-table{table_index:03d}",
        "page_number": page_number,
        "table_index": table_index,
        "bbox": table_bbox,
        "row_count": row_count,
        "column_count": column_count,
        "cells": cells,
        "confidence": confidence,
        "unresolved_positions": unresolved_positions,
        "vector_edge_count": vector_edge_count,
    }


def _cluster_pdf_coordinates(
    values, tolerance: float = _PDF_TABLE_COORDINATE_TOLERANCE
) -> list[float]:
    groups = []
    for value in sorted(float(item) for item in values):
        if not groups or abs(value - groups[-1][0] / groups[-1][1]) > tolerance:
            groups.append([value, 1])
        else:
            groups[-1][0] += value
            groups[-1][1] += 1
    return [total / count for total, count in groups]


def _pdf_coordinate_index(boundaries: list[float], value: float) -> int | None:
    if not boundaries:
        return None
    insertion = bisect_left(boundaries, value)
    candidates = []
    if insertion < len(boundaries):
        candidates.append(insertion)
    if insertion:
        candidates.append(insertion - 1)
    index = min(candidates, key=lambda candidate: abs(boundaries[candidate] - value))
    if abs(boundaries[index] - value) > _PDF_TABLE_COORDINATE_TOLERANCE:
        return None
    return index


def _pdf_table_vector_edge_count(
    drawings: list[dict], table_bbox: tuple[float, ...]
) -> int:
    table_rect = fitz.Rect(table_bbox)
    count = 0
    for drawing in drawings:
        drawing_rect = drawing.get("rect")
        if drawing_rect is not None and not fitz.Rect(drawing_rect).intersects(
            table_rect
        ):
            continue
        for item in drawing.get("items", []):
            if not item:
                continue
            if item[0] == "l":
                count += 1
            elif item[0] == "re":
                count += 4
    return count


class _PdfTextPresenceIndex:
    """通过页面已有字符边界定位空白区域，保留有文字候选的原生文本回退。"""

    def __init__(self, raw_page):
        self.complete = "blocks" in raw_page
        rectangles = []
        for block in raw_page.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if "chars" not in span:
                        self.complete = False
                    for char in span.get("chars", []):
                        bbox = char.get("bbox")
                        if bbox is None or len(bbox) != 4:
                            self.complete = False
                        else:
                            rectangles.append(tuple(bbox))
        self.rectangles = sorted(rectangles, key=lambda bbox: bbox[1])
        self.starts = [bbox[1] for bbox in self.rectangles]
        self.bottoms = []
        maximum = float("-inf")
        for bbox in self.rectangles:
            maximum = max(maximum, bbox[3])
            self.bottoms.append(maximum)

    def may_have_text(self, bbox):
        if not self.complete:
            return True
        tolerance = _PDF_TABLE_COORDINATE_TOLERANCE
        expanded = (
            bbox[0] - tolerance,
            bbox[1] - tolerance,
            bbox[2] + tolerance,
            bbox[3] + tolerance,
        )
        start = bisect_right(self.bottoms, expanded[1])
        end = bisect_left(self.starts, expanded[3])
        return any(
            _pdf_rects_intersect(self.rectangles[index], expanded)
            for index in range(start, end)
        )


class _PdfNontextIndex:
    """按纵向中心坐标索引图片和有效图形，单元格只检查相邻候选。"""

    def __init__(self, images, drawings):
        self.images = sorted(
            ((bbox[1] + bbox[3]) / 2, (bbox[0] + bbox[2]) / 2)
            for bbox in images
            if max(0.0, bbox[2] - bbox[0]) * max(0.0, bbox[3] - bbox[1]) >= 4
        )
        points = []
        for drawing in drawings:
            value = drawing.get("rect")
            if value is None:
                continue
            x0, y0, x1, y1 = (float(coordinate) for coordinate in value)
            if x1 - x0 >= 2 and y1 - y0 >= 2:
                points.append(((y0 + y1) / 2, (x0 + x1) / 2))
        self.drawings = sorted(points)

    @staticmethod
    def _contains(points, bbox):
        tolerance = _PDF_TABLE_COORDINATE_TOLERANCE
        start = bisect_left(points, (bbox[1] - tolerance, float("-inf")))
        end = bisect_right(points, (bbox[3] + tolerance, float("inf")))
        return any(
            bbox[0] - tolerance <= points[index][1] <= bbox[2] + tolerance
            for index in range(start, end)
        )

    def contains(self, bbox):
        if self._contains(self.images, bbox):
            return True
        x0, y0, x1, y1 = bbox
        inset = max(1.0, min(x1 - x0, y1 - y0) * 0.04)
        interior = (x0 + inset, y0 + inset, x1 - inset, y1 - inset)
        return (
            interior[2] > interior[0]
            and interior[3] > interior[1]
            and self._contains(self.drawings, interior)
        )


def _pdf_page_text_elements(
    raw_page: dict, table_rects: list[tuple[float, ...]]
) -> list[dict]:
    elements = []
    source_order = 0
    for block in raw_page.get("blocks", []):
        if block.get("type", 0) != 0:
            continue
        for line in block.get("lines", []):
            source_order += 1
            _source, text = _pdf_line_text_variants(line)
            text = text.strip()
            bbox = line.get("bbox")
            if not text or not bbox:
                continue
            normalized_bbox = tuple(float(value) for value in bbox)
            if any(
                _pdf_rect_contains_center(table_rect, normalized_bbox)
                for table_rect in table_rects
            ):
                continue
            elements.append(
                {
                    "bbox": normalized_bbox,
                    "content": text,
                    "plain_text": text,
                    "kind": "text",
                    "source_order": source_order,
                }
            )
    return elements


def _pdf_rect_contains_center(
    container_bbox: tuple[float, ...], item_bbox: tuple[float, ...]
) -> bool:
    center_x = (item_bbox[0] + item_bbox[2]) / 2
    center_y = (item_bbox[1] + item_bbox[3]) / 2
    return (
        container_bbox[0] - _PDF_TABLE_COORDINATE_TOLERANCE
        <= center_x
        <= container_bbox[2] + _PDF_TABLE_COORDINATE_TOLERANCE
        and container_bbox[1] - _PDF_TABLE_COORDINATE_TOLERANCE
        <= center_y
        <= container_bbox[3] + _PDF_TABLE_COORDINATE_TOLERANCE
    )


def _pdf_page_element_sort_key(element: dict) -> tuple:
    bbox = element["bbox"]
    return (
        round(float(bbox[1]), 1),
        round(float(bbox[0]), 1),
        1 if element["kind"] == "table" else 0,
        int(element["source_order"]),
    )


def _format_pdf_table_html(model: dict) -> str:
    normalized_cells = {}
    for cell in model["cells"]:
        anchor = _pdf_cell_anchor_reference(cell)
        original_range = _pdf_cell_reference(cell)
        for row_index in range(cell["row"], cell["row"] + cell["rowspan"]):
            for column_index in range(cell["column"], cell["column"] + cell["colspan"]):
                inherited = (row_index, column_index) != (cell["row"], cell["column"])
                normalized_cells[(row_index, column_index)] = {
                    "cell": cell,
                    "anchor": anchor,
                    "original_range": original_range,
                    "inherited": inherited,
                }
    lines = [
        (
            f"[PDF结构化表格 {model['id']} 开始；"
            f"{model['row_count']}行×{model['column_count']}列；"
            f"置信度={model['confidence']}；视图=归一化；"
            f"合并继承位置通过 data-inherited-from 标记]"
        ),
        (
            f'<table id="{model["id"]}" data-confidence="{model["confidence"]}" '
            'data-view="normalized">'
        ),
    ]
    for row_index in range(model["row_count"]):
        lines.append("  <tr>")
        for column_index in range(model["column_count"]):
            cell_ref = _pdf_position_reference(row_index, column_index)
            normalized = normalized_cells.get((row_index, column_index))
            if normalized is None:
                lines.append(
                    f'    <td data-cell="{cell_ref}" data-unresolved="true">'
                    "[未能可靠还原的单元格]</td>"
                )
                continue
            cell = normalized["cell"]
            attributes = [f'data-cell="{cell_ref}"']
            if cell["rowspan"] > 1 or cell["colspan"] > 1:
                attributes.append(
                    f'data-original-range="{normalized["original_range"]}"'
                )
            if normalized["inherited"]:
                attributes.append(f'data-inherited-from="{normalized["anchor"]}"')
                if cell["text"]:
                    value = html.escape(cell["text"])
                elif cell["has_nontext_content"]:
                    value = f"[合并覆盖，继承自{normalized['anchor']}的非文本内容]"
                else:
                    value = f"[合并覆盖，继承自{normalized['anchor']}]"
            else:
                if cell["rowspan"] > 1:
                    attributes.append(f'data-original-rowspan="{cell["rowspan"]}"')
                if cell["colspan"] > 1:
                    attributes.append(f'data-original-colspan="{cell["colspan"]}"')
                value = html.escape(cell["text"])
                if not value:
                    if cell["has_nontext_content"]:
                        attributes.append('data-non-text="true"')
                        value = "[非文本图形或图标]"
                    else:
                        attributes.append('data-empty="true"')
                        value = "[空单元格]"
            lines.append(f"    <td {' '.join(attributes)}>{value}</td>")
        lines.append("  </tr>")
    if model["unresolved_positions"]:
        positions = ",".join(
            f"{get_column_letter(column + 1)}{row + 1}"
            for row, column in model["unresolved_positions"]
        )
        lines.append(f"  <!-- 未能可靠还原的网格位置：{positions} -->")
    lines.extend(
        [
            "</table>",
            f"[PDF结构化表格 {model['id']} 结束]",
        ]
    )
    return "\n".join(lines)


def _pdf_cell_reference(cell: dict) -> str:
    start = _pdf_cell_anchor_reference(cell)
    end = _pdf_position_reference(
        cell["row"] + cell["rowspan"] - 1,
        cell["column"] + cell["colspan"] - 1,
    )
    return start if start == end else f"{start}:{end}"


def _pdf_cell_anchor_reference(cell: dict) -> str:
    return _pdf_position_reference(cell["row"], cell["column"])


def _pdf_position_reference(row: int, column: int) -> str:
    return f"{get_column_letter(column + 1)}{row + 1}"


def _extract_pymupdf_page_hyperlinks(page) -> list[tuple[str, str, bool | None]]:
    hyperlinks = []
    try:
        links = page.get_links()
    except Exception:
        return hyperlinks
    for link in links:
        target = _clean_hyperlink_target(link.get("uri") or link.get("file"))
        internal_target_exists = None
        if not target and link.get("kind") == fitz.LINK_GOTO:
            page_index = link.get("page")
            if isinstance(page_index, int):
                target = f"#page={page_index + 1}"
                internal_target_exists = 0 <= page_index < page.parent.page_count
        if not target and link.get("nameddest"):
            target = f"#{_clean_hyperlink_target(link.get('nameddest')).lstrip('#')}"
        label = ""
        link_rect = link.get("from")
        if link_rect:
            try:
                label = page.get_textbox(link_rect).strip()
            except Exception:
                pass
        item = (label, target, internal_target_exists)
        if item not in hyperlinks:
            hyperlinks.append(item)
    return hyperlinks


def _extract_pypdf_page_hyperlinks(page) -> list[tuple[str, str, bool | None]]:
    hyperlinks = []
    try:
        annotations = page.get("/Annots") or []
    except Exception:
        return hyperlinks
    for annotation_ref in annotations:
        try:
            annotation = annotation_ref.get_object()
            if annotation.get("/Subtype") != "/Link":
                continue
            action = annotation.get("/A") or {}
            target = _clean_hyperlink_target(action.get("/URI"))
        except Exception:
            continue
        item = ("", target, None)
        if target and item not in hyperlinks:
            hyperlinks.append(item)
    return hyperlinks


def _append_hyperlinks(
    text: str,
    hyperlinks: list[tuple[str, str, bool | None]],
) -> str:
    annotations = [
        f"[超链接] {_format_hyperlink_text(label, target)}"
        for label, target, _internal_target_exists in hyperlinks
    ]
    return "\n".join(part for part in (str(text or "").strip(), *annotations) if part)


def _select_pdf_page_text(pypdf_text: str, pymupdf_text: str) -> str:
    pypdf_text = str(pypdf_text or "")
    pymupdf_text = str(pymupdf_text or "")
    if not pymupdf_text.strip():
        return pypdf_text
    if not pypdf_text.strip():
        return pymupdf_text

    pypdf_quality = _pdf_text_quality(pypdf_text)
    pymupdf_quality = _pdf_text_quality(pymupdf_text)
    pypdf_useful = pypdf_quality["useful_chars"]
    pymupdf_useful = pymupdf_quality["useful_chars"]

    if (
        pymupdf_quality["bad_char_ratio"] >= 0.01
        and pypdf_quality["bad_char_ratio"] <= pymupdf_quality["bad_char_ratio"] * 0.5
        and pypdf_useful >= pymupdf_useful * 0.7
    ):
        return pypdf_text

    if pypdf_quality["bad_char_ratio"] > max(
        0.01, pymupdf_quality["bad_char_ratio"] * 2
    ):
        return pymupdf_text

    pymupdf_tokens = pymupdf_quality["tokens"]
    pypdf_tokens = pypdf_quality["tokens"]
    pymupdf_covered_by_pypdf = _shared_pdf_token_ratio(pymupdf_tokens, pypdf_tokens)
    pypdf_covered_by_pymupdf = _shared_pdf_token_ratio(pypdf_tokens, pymupdf_tokens)
    if (
        pymupdf_tokens
        and pypdf_tokens
        and pymupdf_covered_by_pypdf < 0.3
        and pypdf_covered_by_pymupdf < 0.3
    ):
        return pypdf_text

    extra_pypdf_chars = pypdf_useful - pymupdf_useful
    pypdf_relative_gain = pypdf_useful / max(pymupdf_useful, 1)
    pypdf_is_materially_more_complete = extra_pypdf_chars >= max(
        20, int(pymupdf_useful * 0.15)
    )
    if (
        pypdf_is_materially_more_complete
        and pypdf_relative_gain >= 1.15
        and pymupdf_covered_by_pypdf >= 0.8
    ):
        return pypdf_text
    return pymupdf_text


def _pdf_text_quality(text: str) -> dict:
    visible_chars = [char for char in text if not char.isspace()]
    bad_chars = [char for char in visible_chars if _is_bad_pdf_text_char(char)]
    useful_chars = len(visible_chars) - len(bad_chars)
    tokens = set(re.findall(r"\w+", text.casefold(), flags=re.UNICODE))
    return {
        "useful_chars": useful_chars,
        "bad_char_ratio": len(bad_chars) / max(len(visible_chars), 1),
        "tokens": tokens,
    }


def _is_bad_pdf_text_char(char: str) -> bool:
    if char == "\ufffd":
        return True
    return unicodedata.category(char) in {"Cc", "Cs", "Co", "Cn"}


def _shared_pdf_token_ratio(
    primary_tokens: set[str], candidate_tokens: set[str]
) -> float:
    if not primary_tokens:
        return 1.0
    return len(primary_tokens & candidate_tokens) / len(primary_tokens)


def _normalize_pdf_overlapping_spaces(text: str, raw_page: dict) -> str:
    """Remove encoded spaces that have no visible advance on the PDF page."""
    if not text or not raw_page:
        return text

    source_counts = {}
    replacements = {}
    for block in raw_page.get("blocks", []):
        for line in block.get("lines", []):
            source, normalized = _pdf_line_text_variants(line)
            if not source:
                continue
            occurrence = source_counts.get(source, 0) + 1
            source_counts[source] = occurrence
            if normalized != source:
                replacements[(source, occurrence)] = normalized

    if not replacements:
        return text

    text_counts = {}
    normalized_lines = []
    for line in text.splitlines(keepends=True):
        content = line.rstrip("\r\n")
        line_ending = line[len(content) :]
        occurrence = text_counts.get(content, 0) + 1
        text_counts[content] = occurrence
        normalized_lines.append(
            replacements.get((content, occurrence), content) + line_ending
        )
    return "".join(normalized_lines)


def _pdf_line_text_variants(line: dict) -> tuple[str, str]:
    cached = line.get("_text_variants")
    if isinstance(cached, tuple) and len(cached) == 2:
        return cached
    chars = []
    for span in line.get("spans", []):
        size = float(span.get("size") or 0)
        for char in span.get("chars", []):
            value = str(char.get("c") or "")
            if value:
                chars.append(
                    {"value": value, "origin": char.get("origin"), "size": size}
                )

    source = "".join(char["value"] for char in chars)
    if " " not in source or len(chars) < 3:
        result = (source, source)
        line["_text_variants"] = result
        return result

    direction = line.get("dir") or (1.0, 0.0)
    try:
        direction_x = float(direction[0])
        direction_y = float(direction[1])
        direction_length = (direction_x**2 + direction_y**2) ** 0.5
    except (TypeError, ValueError, IndexError):
        result = (source, source)
        line["_text_variants"] = result
        return result
    if direction_length <= 0:
        result = (source, source)
        line["_text_variants"] = result
        return result
    direction_x /= direction_length
    direction_y /= direction_length

    next_content = [None] * len(chars)
    following = None
    for index in range(len(chars) - 1, -1, -1):
        next_content[index] = following
        if chars[index]["value"].strip():
            following = chars[index]

    remove_indexes = set()
    has_previous_content = False
    for index, char in enumerate(chars):
        if char["value"].strip():
            has_previous_content = True
        if char["value"] != " " or not has_previous_content:
            continue
        next_char = next_content[index]
        if next_char is None:
            continue
        origin = char.get("origin")
        next_origin = next_char.get("origin")
        if not origin or not next_origin:
            continue
        try:
            advance = (float(next_origin[0]) - float(origin[0])) * direction_x + (
                float(next_origin[1]) - float(origin[1])
            ) * direction_y
        except (TypeError, ValueError, IndexError):
            continue
        tolerance = max(0.25, float(char.get("size") or 0) * 0.03)
        if abs(advance) <= tolerance:
            remove_indexes.add(index)

    normalized = "".join(
        char["value"] for index, char in enumerate(chars) if index not in remove_indexes
    )
    result = (source, normalized)
    line["_text_variants"] = result
    return result
